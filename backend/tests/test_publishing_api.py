"""Published files, the importer, and the shape of the API."""
import io
import re
import zipfile
from pathlib import Path

import docx

from app import originality, publisher, store
from app.main import API_VERSION, app
from app.models import GenerateRequest
from app.writer import RUN_OPTIONS
from conftest import DATA, ROOT, make_project


def article(with_appendix=False):
    art = publisher.assemble(make_project())
    if with_appendix:
        B = publisher.Block
        art.sections.append(publisher.Section("appendices", "Appendices", "", [B("subhead", "Appendix A. Data provenance"), B("p", "Twelve documents."),
                                                                              B("subhead", "Appendix B. Variable dictionary"), B("p", "Listed here.")]))
    return art


def test_docx_contents_lists_every_heading_and_links_to_it():
    art = article(with_appendix=True)
    for key, tpl in publisher.TEMPLATES.items():
        path = publisher.to_docx(art, DATA / f"{key}.docx", key, "emerald")
        xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf8")
        settings_xml = zipfile.ZipFile(path).read("word/settings.xml").decode("utf8")
        d = docx.Document(str(path))
        expected = [t for _, t in publisher.contents_entries(art, tpl)]
        assert [q.text for q in d.paragraphs if q.style.name in ("Heading 1", "Heading 2")] == expected, key
        if not tpl.toc:
            assert "TOC \\o" not in xml and "updateFields" not in settings_xml, key
            continue
        texts = [q.text for q in d.paragraphs]
        start = texts.index("Contents")
        assert texts[start + 1:start + 1 + len(expected)] == expected and len(expected) == 6, key
        anchors, marks = re.findall(r'w:anchor="(_Toc\d+)"', xml), re.findall(r'<w:bookmarkStart[^>]*w:name="(_Toc\d+)"', xml)
        assert anchors == marks and len(anchors) == len(expected), key
        assert xml.count('w:fldCharType="begin"') == xml.count('w:fldCharType="end"') and "updateFields" in settings_xml, key


def test_markdown_contents_and_reimport_round_trip():
    p = make_project()
    art = publisher.assemble(p)
    md = publisher.to_markdown(art)
    links = re.findall(r"^\s*- \[(.+?)\]\(#(.+?)\)$", md, re.M)
    assert [t for t, _ in links] == ["1. Introduction", "2. Discussion", "References"] and links[0][1] == "1-introduction"
    for name, write in (("a.md", lambda f: f.write_text(md, encoding="utf8")), ("a.docx", lambda f: publisher.to_docx(art, f, "modern_report", "violet"))):
        write(DATA / name)
        got = publisher.import_document(p, DATA / name)
        assert {"introduction", "discussion", "references"} <= set(got), name
        assert "complement managerial skill" in got["discussion"] and "](#" not in "".join(got.values()), name
        assert len(got["introduction"]) > 300, name                     # the section, not its line in the contents list


def test_destyle_drops_stock_connectors():
    s = originality.destyle("The effect is large. Moreover, the sample was small.\n\nIn conclusion, more work is needed.")
    assert s == "The effect is large. The sample was small.\n\nMore work is needed."


def test_statistics_keep_their_spacing():
    """"p = .05" must never become "p =.05": a space before punctuation is only removed when no digit follows."""
    from app import citations, factcheck
    p = make_project()
    text = "The effect was small (r = .21, p = .041) [R1] . It held in both samples , as expected [R2]."
    assert citations.render_in_text(text, p, ["R1", "R2"]) == "The effect was small (r = .21, p = .041) (Smith, 2021). It held in both samples, as expected (Garcia, 2019)."
    assert originality.destyle("A value of p = .05 was used , and r = .30 .") == "A value of p = .05 was used, and r = .30."
    assert factcheck.sanitise("Scores rose (d = .44) , on average .", set())[0] == "Scores rose (d = .44), on average."


def test_resume_keeps_the_options_the_run_started_with():
    stored = {"web_research": False, "data_analysis": True, "depth": "quick", "agents": 2, "cloud": False, "hallucination_filter": False, "style_pass": False}
    def resolve(req):
        return req.model_copy(update={k: stored[k] for k in RUN_OPTIONS if k not in req.model_fields_set and k in stored})
    assert all(getattr(resolve(GenerateRequest(segments=["introduction"], resume=True)), k) == stored[k] for k in RUN_OPTIONS)
    assert resolve(GenerateRequest(segments=["introduction"], resume=True, style_pass=True)).style_pass is True


def test_interface_and_api_versions_match():
    expected = re.search(r'EXPECTED_API = "([\d.]+)"', (ROOT / "frontend/src/App.jsx").read_text(encoding="utf8")).group(1)
    assert expected == API_VERSION


def test_every_route_the_interface_calls_exists(client):
    paths = {(m.upper(), path) for path, methods in app.openapi()["paths"].items() for m in methods}
    api_js = (ROOT / "frontend/src/api.js").read_text(encoding="utf8")
    # routes that existed before the 1.5 split into modules must all still be there
    for method, path in [("GET", "/api/health"), ("GET", "/api/segments"), ("GET", "/api/tones"), ("GET", "/api/length-profile"), ("GET", "/api/publish/styles"),
                         ("GET", "/api/projects"), ("POST", "/api/projects"), ("GET", "/api/projects/{pid}"), ("PATCH", "/api/projects/{pid}"),
                         ("DELETE", "/api/projects/{pid}"), ("POST", "/api/projects/{pid}/stage/{stage}"), ("POST", "/api/projects/{pid}/references"),
                         ("DELETE", "/api/projects/{pid}/references/{rid}"), ("POST", "/api/projects/{pid}/datasets"), ("DELETE", "/api/projects/{pid}/datasets/{name}"),
                         ("PUT", "/api/projects/{pid}/lengths"), ("PUT", "/api/projects/{pid}/tone"), ("PUT", "/api/projects/{pid}/publish-style"),
                         ("POST", "/api/projects/{pid}/generate"), ("GET", "/api/projects/{pid}/events"), ("POST", "/api/projects/{pid}/review"),
                         ("POST", "/api/projects/{pid}/segments/{key}/revise"), ("POST", "/api/projects/{pid}/segments/{key}/tool"),
                         ("POST", "/api/projects/{pid}/segments/{key}/regenerate"), ("PUT", "/api/projects/{pid}/segments/{key}"),
                         ("POST", "/api/projects/{pid}/segments/{key}/approve"), ("POST", "/api/projects/{pid}/segments/{key}/restore/{index}"),
                         ("GET", "/api/projects/{pid}/files/{path}"),("GET", "/api/projects/{pid}/export/{fmt}"), ("POST", "/api/projects/{pid}/import"),
                         ("GET", "/api/settings/llm"), ("PUT", "/api/settings/llm"), ("DELETE", "/api/settings/llm"), ("POST", "/api/settings/llm/benchmark")]:
        assert (method, path) in paths, (method, path)
    # and every path written in api.js resolves to a route
    known = [re.sub(r"\{[^}]+\}", "[^/]+", p) + "$" for _, p in paths]
    for raw in set(re.findall(r"[`\"](/(?:projects|settings|segments|tones|health|publish|length-profile)[^`\"?]*)", api_js)):
        url = "/api" + re.sub(r"\$\{[^}]+\}", "x", raw)
        assert any(re.match(k, url) for k in known), raw
    assert client.get("/api/health").json()["api_version"] == API_VERSION
    assert client.get("/api/projects/nope-not-an-id").status_code == 404
    assert client.get(f"/api/projects/{make_project().id}/files/../../../llm_settings.json").status_code == 404
    assert not Path(store.project_dir("x" * 12).parent / "nope-not-an-id").exists()
