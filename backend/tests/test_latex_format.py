"""1.7 — LaTeX / BibTeX / RIS export, and "Format an existing article"."""
import io
import re
import zipfile

import docx
import pytest
from docx.shared import Pt
from PIL import Image

from app import citations, formatter, publisher, store
from conftest import DATA, make_project


# ------------------------------------------------------------------ helpers
def check_latex(tex: str, files: set[str], bib: str = "") -> None:
    """No LaTeX engine is needed for these: they catch what would stop a compile."""
    body = re.sub(r"\\[\\{}%&#_$]", "", tex)                       # escaped characters do not count
    body = "\n".join(l for l in body.splitlines() if not l.lstrip().startswith("%"))
    assert body.count("{") == body.count("}"), "unbalanced braces"
    depth = 0
    for ch in body:
        depth += ch == "{"
        depth -= ch == "}"
        assert depth >= 0, "a closing brace before its opening one"
    begins, ends = re.findall(r"\\begin\{(\w+\*?)\}", tex), re.findall(r"\\end\{(\w+\*?)\}", tex)
    assert sorted(begins) == sorted(ends), (begins, ends)
    assert tex.count("\\begin{document}") == 1 and tex.rstrip().endswith("\\end{document}")
    for line in body.split("\\begin{document}")[1].splitlines():
        stripped = re.sub(r"\$[^$]*\$", "", line)                   # inline maths RAF wrote itself
        if "tabularx" in line or " & " in line or line.startswith("\\"):
            continue
        assert not re.search(r"[&#%_^~]", stripped) or "\\parencite" in stripped and not re.search(r"[&#%^]", re.sub(r"\\parencite\{[^}]*\}", "", stripped)), line
    cited = {k for group in re.findall(r"\\parencite\{([^}]*)\}", tex) for k in group.split(",")}
    assert cited <= set(re.findall(r"@\w+\{([^,]+),", bib)), "a citation without a .bib entry"
    for name in re.findall(r"\\includegraphics\[[^\]]*\]\{([^}]+)\}", tex):
        assert name in files, f"{name} is not in the zip"


def unzip(blob: bytes) -> dict[str, bytes]:
    z = zipfile.ZipFile(io.BytesIO(blob))
    return {n: z.read(n) for n in z.namelist()}


def sample_docx(path):
    d = docx.Document()
    d.add_paragraph("Water Pricing and Household Demand", style="Title")
    d.add_paragraph("Asha Verma, Tom Okoro")
    d.add_paragraph("Department of Economics, Example University")
    d.add_heading("Abstract", level=1)
    d.add_paragraph("This study estimates how households respond to block tariffs. Demand falls when prices rise, but less than planners assume.")
    d.add_paragraph("Keywords: water pricing; demand; tariffs")
    d.add_heading("1. Introduction", level=1)
    par = d.add_paragraph("Tariff reform is ")
    par.add_run("widely").bold = True
    par.add_run(" debated, and ")
    par.add_run("ex ante").italic = True
    par.add_run(" estimates differ (Rao, 2019). Use is measured in m")
    par.add_run("3").font.superscript = True
    par.add_run(" per month; * marks p < .05 and ** marks p < .01.")
    d.add_heading("2. METHODS", level=1)
    d.add_heading("2.1 Data", level=2)
    d.add_paragraph("Billing records for 4,200 households over 36 months were used (Lee & Park, 2021; Ncube et al., 2020).")
    d.add_paragraph("Table 1. Sample by tariff block")
    t = d.add_table(rows=3, cols=3)
    for r, row in enumerate([["Block", "Households", "Share %"], ["Low", "2,100", "50"], ["High", "2,100", "50"]]):
        for c, text in enumerate(row):
            t.cell(r, c).text = text
    buf = io.BytesIO()
    Image.new("RGB", (240, 120), (30, 110, 200)).save(buf, "PNG")
    buf.seek(0)
    d.add_picture(buf)
    d.add_paragraph("Figure 1. Monthly use by block")
    d.add_heading("Acknowledgements", level=1)
    d.add_paragraph("We thank the utility for the data.")
    d.add_heading("References", level=1)
    for ref in ("Lee, S., & Park, J. (2021). Block tariffs and demand. Water Economics, 12(3), 44–61.",
                "Rao, P. (2019). Pricing water. Journal of Utilities, 7(1), 1–20.",
                "Smith, A. (2015). An entry nobody cites. Old Journal, 1, 1–2."):
        d.add_paragraph(ref)
    d.styles["Normal"].font.size = Pt(11)
    d.save(path)
    return path


def upload(client, path, name=None):
    with open(path, "rb") as fh:
        return client.post("/api/format", files={"file": (name or path.name, fh.read(), "application/octet-stream")})


# ------------------------------------------------------------------ LaTeX, BibTeX, RIS
def test_latex_export_in_every_template(client, monkeypatch):
    p = make_project(approved=True)
    p.authors, p.affiliation = ["Jane Doe", "John Smith"], "Example University"
    p.segments["discussion"].content += "\n\nCosts fell by 12% & output rose (ρ = .41, p < .05) in the 2019–2021 file_set #2 [R2, R1]."
    store.save(p)
    monkeypatch.setattr(citations, "details", lambda doi: {"volume": "9", "issue": "2", "pages": "101-120"})
    for template in publisher.TEMPLATES:
        r = client.get(f"/api/projects/{p.id}/export/tex?template={template}&palette=emerald")
        assert r.status_code == 200 and r.headers["content-type"] == "application/zip", template
        files = unzip(r.content)
        assert set(files) == {"main.tex", "references.bib", "README.txt"}, template
        tex, bib = files["main.tex"].decode("utf8"), files["references.bib"].decode("utf8")
        check_latex(tex, set(files), bib)
        assert "\\section{Introduction}" in tex and "\\section{Discussion}" in tex and "\\printbibliography[title={References}]" in tex and "\\nocite{*}" in tex
        assert "lags behind larger competitors \\parencite{smith2021digital}." in tex and "\\parencite{garcia2019managers,smith2021digital}" in tex
        assert "12\\% \\& output rose ($\\rho$ = .41, p \\textless{} .05) in the 2019--2021 file\\_set \\#2" in tex
        assert "[R1]" not in tex and "(Smith, 2021)" not in tex                                   # real citations, not text
        accent = "065F46" if publisher.TEMPLATES[template].color else "000000"                 # the APA manuscript is monochrome
        assert "\\usepackage[style=apa,backend=biber]{biblatex}" in tex and "\\definecolor{accent}{HTML}{%s}" % accent in tex
    assert "\\documentclass[man,12pt,a4paper]{apa7}" in unzip(client.get(f"/api/projects/{p.id}/export/tex?template=apa_manuscript").content)["main.tex"].decode()
    ieee = unzip(client.get(f"/api/projects/{p.id}/export/tex?template=conference").content)["main.tex"].decode()
    assert "{IEEEtran}" in ieee and "\\begin{IEEEkeywords}" in ieee and "\\IEEEauthorblockN{Jane Doe, John Smith}" in ieee
    modern = unzip(client.get(f"/api/projects/{p.id}/export/tex?template=modern_report").content)["main.tex"].decode()
    assert "\\tableofcontents" in modern and "\\colorbox{accentfill}" in modern and "Digital Adoption and Productivity in Small Firms" in modern
    assert "Evidence from Emerging Economies" in modern                                           # the subtitle


def test_bibtex_and_ris(client, monkeypatch):
    p = make_project()
    p.references[0].doi = "10.1000/xyz123"
    p.references.append(p.references[0].model_copy(update={"id": "dup000000001", "filename": "b.pdf", "title": "Digital adoption revisited"}))
    p.options["web_research"] = True
    store.save(p)
    monkeypatch.setattr(citations, "details", lambda doi: {"volume": "9", "issue": "2", "pages": "101-120"})
    bib = client.get(f"/api/projects/{p.id}/export/bib")
    assert bib.status_code == 200 and 'filename="' in bib.headers["content-disposition"]
    assert "@article{smith2021digital,\n  author = {Smith, J.},\n  title = {{Digital adoption in small firms}},\n  journal = {Journal of Small Business},\n  year = {2021}" in bib.text
    assert "volume = {9}" in bib.text and "number = {2}" in bib.text and "pages = {101--120}" in bib.text and "doi = {10.1000/xyz123}" in bib.text
    assert "@article{smith2021digitalb," in bib.text and "@article{garcia2019managers," in bib.text       # same author and year: keys stay unique
    ris = client.get(f"/api/projects/{p.id}/export/ris").text
    assert ris.count("TY  - JOUR") == 3 and "AU  - Garcia, M." in ris and "DO  - 10.1000/xyz123" in ris and ris.count("ER  - ") == 3
    assert "Water Economics, 12(3), 44–61." in citations.format_reference(formatter.Work("T", ["Lee, S."], "2021", "Water Economics", "", "12", "3", "44–61"), "APA")


# ------------------------------------------------------------------ Format an existing article
def test_docx_is_brought_in_with_its_structure(client):
    out = upload(client, sample_docx(DATA / "water.docx")).json()
    found, p = out["found"], out["project"]
    assert found == {"title": "Water Pricing and Household Demand", "authors": ["Asha Verma", "Tom Okoro"], "abstract": True, "keywords": True,
                     "sections": ["Introduction", "Methods", "Acknowledgements"], "references": 3, "tables": 1, "figures": 1, "lossy": False}
    assert p["mode"] == "format" and p["stage"] == "studio" and p["affiliation"] == "Department of Economics, Example University"
    rows = [(r["title"], r["kind"], r["numbered"], r["status"]) for r in p["structure"]]
    assert rows == [("Title", "builtin", False, "approved"), ("Abstract", "builtin", False, "approved"), ("Keywords", "builtin", False, "approved"),
                    ("Introduction", "manual", True, "approved"), ("Methods", "manual", True, "approved"), ("Acknowledgements", "manual", False, "approved"),
                    ("References", "builtin", False, "approved")]
    segs = {s["title"]: s for s in p["segments"].values()}
    assert segs["Keywords"]["content"] == "water pricing; demand; tariffs" and "Keywords" not in segs["Abstract"]["content"]
    intro = segs["Introduction"]["content"]
    assert "is **widely** debated, and *ex ante* estimates" in intro and "in m^3^ per month" in intro
    assert r"\* marks p < .05 and \*\* marks p < .01" in intro                                   # the document's own asterisks stay asterisks
    methods = segs["Methods"]
    assert methods["content"].startswith("## Data\n\nBilling records") and "[[TABLE 1]]" in methods["content"] and "[[FIGURE 1]]" in methods["content"]
    assert "Table 1." not in methods["content"] and "Figure 1." not in methods["content"]          # captions moved onto the table and figure
    assert methods["tables"][0]["caption"] == "Sample by tariff block" and methods["tables"][0]["rows"][1] == ["High", "2,100", "50"]
    assert methods["figures"][0]["caption"] == "Monthly use by block" and (DATA / "projects" / p["id"] / "figures" / "figure_1.png").exists()

    pid = p["id"]
    md = client.get(f"/api/projects/{pid}/export/md").text                                        # already approved: publishes at once
    assert "## 1. Introduction" in md and "## 2. Methods" in md and "### Data" in md and "## Acknowledgements" in md
    assert "**widely**" in md and "m<sup>3</sup>" in md and "* marks p < .05 and ** marks p < .01" in md and "(Rao, 2019)" in md
    assert "| Low | 2,100 | 50 |" in md and "- Rao, P. (2019). Pricing water." in md
    word = docx.Document(io.BytesIO(client.get(f"/api/projects/{pid}/export/docx?template=classic_journal").content))
    runs = [r for par in word.paragraphs for r in par.runs]
    assert any(r.bold and r.text == "widely" for r in runs) and any(r.italic and r.text == "ex ante" for r in runs) and any(r.font.superscript and r.text == "3" for r in runs)
    assert [par.text for par in word.paragraphs if par.style.name == "Heading 2"] == ["Data"] and len(word.tables) == 1 and len(word.inline_shapes) == 1
    assert client.get(f"/api/projects/{pid}/export/pdf").content[:4] == b"%PDF"
    files = unzip(client.get(f"/api/projects/{pid}/export/tex").content)
    tex = files["main.tex"].decode("utf8")
    check_latex(tex, set(files))
    assert set(files) == {"main.tex", "README.txt", "figures/figure_1.png"}                         # no .bib: the list is the author's own text
    assert "is \\textbf{widely} debated, and \\textit{ex ante} estimates" in tex and "m\\textsuperscript{3}" in tex and "\\subsection*{Data}" in tex
    assert "\\item Rao, P. (2019). Pricing water." in tex and "biblatex" not in tex and "\\section*{Acknowledgements}" in tex
    assert client.get(f"/api/projects/{pid}/export/bib").status_code == 422
    assert client.post(f"/api/projects/{pid}/segments/{[k for k, s in p['segments'].items() if s['title'] == 'Methods'][0]}/regenerate").status_code == 422


def test_markdown_text_and_pdf_are_read(client):
    (DATA / "paper.md").write_text("# Night Markets\n\n## Abstract\n\nMarkets at night *change* how streets are used.\n\n**Keywords:** markets, streets\n\n"
                                   "## 1. Introduction\n\nFirst paragraph.\n\nSecond **bold** paragraph.\n\n### Scope\n\nThree districts.\n\n"
                                   "| City | Stalls |\n|---|---|\n| Pune | 120 |\n| Hue | 80 |\n\n## References\n\n- Tran, L. (2020). Street trade. Urban Studies.\n", encoding="utf8")
    out = upload(client, DATA / "paper.md").json()
    assert out["found"]["title"] == "Night Markets" and out["found"]["sections"] == ["Introduction"] and out["found"]["tables"] == 1 and out["found"]["references"] == 1
    segs = {s["title"]: s for s in out["project"]["segments"].values()}
    assert segs["Abstract"]["content"] == "Markets at night *change* how streets are used." and segs["Keywords"]["content"] == "markets; streets"
    assert segs["Introduction"]["content"].split("\n\n") == ["First paragraph.", "Second **bold** paragraph.", "## Scope", "Three districts.", "[[TABLE 1]]"]

    (DATA / "plain.txt").write_text("A Plain Report\n\nINTRODUCTION\n\nSome text about costs of 5 * 3 units.\n\n2. Results\n\nMore text here.\n", encoding="utf8")
    plain = upload(client, DATA / "plain.txt").json()
    assert plain["found"]["sections"] == ["Introduction", "Results"]
    assert r"5 \* 3 units" in next(s["content"] for s in plain["project"]["segments"].values() if s["title"] == "Introduction")

    pdf = publisher.to_pdf(publisher.assemble(make_project()), DATA / "made.pdf", "classic_journal", "violet")
    back = upload(client, pdf).json()["found"]
    assert back["lossy"] is True and {"Introduction", "Discussion"} <= set(back["sections"]) and back["title"].startswith("Digital Adoption and Productivity")
    assert upload(client, DATA / "plain.txt", "notes.rtf").status_code == 422
    (DATA / "empty.md").write_text("# Only a title\n", encoding="utf8")
    assert upload(client, DATA / "empty.md").status_code == 422


def test_checks_and_reference_restyle(client, monkeypatch):
    pid = upload(client, sample_docx(DATA / "water2.docx")).json()["project"]["id"]
    out = client.get(f"/api/projects/{pid}/format/checks").json()
    refs = out["references"]
    assert refs["style"] == "author–date" and refs["references"] == 3 and refs["missing"] == ["Ncube (2020)"]      # cited, not in the list
    assert len(refs["uncited"]) == 1 and refs["uncited"][0].startswith("Smith, A. (2015)")                             # listed, never cited
    assert [s["title"] for s in out["style"]] == [] or all("score" in s for s in out["style"])

    works = {"Lee": formatter.Work("Block tariffs and demand", ["Lee, S.", "Park, J."], "2021", "Water Economics", "10.1/a", "12", "3", "44-61"),
             "Rao": formatter.Work("Pricing water", ["Rao, P."], "2019", "Journal of Utilities", "", "7", "1", "1-20")}
    monkeypatch.setattr(formatter, "lookup", lambda entry: works.get(entry.split(",")[0]))
    assert client.post(f"/api/projects/{pid}/format/restyle", json={"style": "IEEE"}).status_code == 422               # would need the text renumbered
    done = client.post(f"/api/projects/{pid}/format/restyle", json={"style": "Harvard"}).json()
    assert done["restyled"] == 2 and done["total"] == 3 and len(done["unchanged"]) == 1 and done["project"]["citation_style"] == "Harvard"
    lines = done["project"]["segments"]["references"]["content"].splitlines()
    assert lines[0] == "Lee, S. and Park, J. (2021) 'Block tariffs and demand', Water Economics, 12(3), 44-61. Available at: https://doi.org/10.1/a"
    assert lines[2].startswith("Smith, A. (2015). An entry nobody cites.")                                             # not identified: left exactly as it was
    client.put(f"/api/projects/{pid}/segments/{next(k for k, s in done['project']['segments'].items() if s['title'] == 'Introduction')}", json={"content": "Edited."})
    assert store.load(pid).segments["references"].content.splitlines() == lines                                        # and never rebuilt behind the author's back

    numbered = make_project()
    store.mutate(numbered.id, lambda p: (setattr(p, "mode", "format"), setattr(p.segments["references"], "content", "[1] A. One, 2019.\n[2] B. Two, 2020.\n[3] C. Three, 2021."),
                                         setattr(p.segments["discussion"], "content", "As shown in [1] and later in [3–4].")))
    check = formatter.reference_check(store.load(numbered.id))
    assert check["style"] == "numbered" and check["missing"] == ["[4]"] and check["uncited"] == ["[2]"]
    assert client.get(f"/api/projects/{make_project().id}/format/checks").status_code == 422                           # only for formatted imports
