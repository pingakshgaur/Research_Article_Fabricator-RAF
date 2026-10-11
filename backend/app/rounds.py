"""Revision rounds: turn the comments of a reviewer or supervisor into revisions and a response letter.

Stored in  data/projects/<id>/revisions/round_<id>.json

1. analyse — the comments are split into separate points; each point is linked to the segment it concerns and given a
             revision instruction. The author accepts, edits or skips every point before anything is changed.
2. apply   — the article is saved as a revision, the accepted points are applied segment by segment with the ordinary
             revision engine (so every filter still runs), and the result is saved as a second revision.
3. letter  — a point-by-point "Response to reviewers" is drafted from what actually changed between the two revisions.
"""
from __future__ import annotations

import json
import re
import time
import uuid
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from . import history, llm, revisions, store, writer
from .ingest import split_sentences
from .models import Project

EDITABLE_SKIP = {"references", "appendices"}
HINTS = (
    ("title", ("title",)), ("abstract", ("abstract",)), ("keywords", ("keyword",)),
    ("limitations", ("limitation", "future research", "future work", "generalis", "generaliz")),
    ("methodology", ("method", "sample", "sampling", "data collection", "research design", "procedure", "validity", "reliability", "ethic")),
    ("results", ("result", "table", "figure", "statistic", "regression", "correlation", "p-value", "effect size", "finding")),
    ("literature_review", ("literature", "theor", "framework", "prior work", "previous studies", "hypothes")),
    ("discussion", ("discussion", "interpret", "explain", "causal", "mechanism")),
    ("conclusion", ("conclusion", "implication", "recommendation", "contribution")),
    ("introduction", ("introduction", "research question", "research gap", "motivation", "objective", "aim")),
)


def _path(pid: str, rid: str) -> Path:
    d = store.project_dir(pid) / "revisions"
    d.mkdir(exist_ok=True)
    return d / f"round_{rid}.json"


def save(pid: str, rnd: dict) -> dict:
    tmp = _path(pid, rnd["id"]).with_suffix(".tmp")
    tmp.write_text(json.dumps(rnd, ensure_ascii=False, indent=1), encoding="utf-8")
    tmp.replace(_path(pid, rnd["id"]))
    return rnd


def get(pid: str, rid: str) -> dict:
    if not re.fullmatch(r"[a-z0-9]{6,16}", rid) or not _path(pid, rid).exists():
        raise KeyError(rid)
    return json.loads(_path(pid, rid).read_text(encoding="utf-8"))


def list_(pid: str) -> list[dict]:
    out = []
    for f in (store.project_dir(pid) / "revisions").glob("round_*.json"):
        try:
            out.append(json.loads(f.read_text(encoding="utf-8")))
        except json.JSONDecodeError:
            continue
    return sorted(out, key=lambda r: r["created"], reverse=True)


def delete(pid: str, rid: str) -> None:
    get(pid, rid)
    _path(pid, rid).unlink(missing_ok=True)


def editable(p: Project) -> list[str]:
    return [k for k in revisions._order(p, p.segments) if p.segments[k].content and k not in EDITABLE_SKIP]


def create(pid: str, comments: str, name: str = "") -> dict:
    n = len(list_(pid)) + 1
    rnd = {"id": uuid.uuid4().hex[:10], "n": n, "name": name.strip()[:80] or f"Round {n}", "created": time.time(),
           "comments": comments.strip(), "status": "analysing", "points": [], "before": None, "after": None, "letter": "",
           "analysed_by": "", "error": ""}
    history.record(pid, "round.started", f"Revision round started — “{rnd['name']}” ({len(rnd['comments'].split())} words of comments)", round=rnd["id"])
    return save(pid, rnd)


# ---------------------------------------------------------------------------- 1. analyse
def guess_segment(text: str, keys: list[str]) -> str:
    low = text.lower()
    for key, words in HINTS:
        if key in keys and any(w in low for w in words):
            return key
    return next((k for k in ("discussion", "introduction") if k in keys), keys[0])


def split_plain(comments: str, keys: list[str]) -> list[dict]:
    """Split comments without the model: one point per numbered item, bullet or paragraph."""
    text = re.sub(r"\r\n?", "\n", comments)
    parts = re.split(r"\n\s*\n|\n(?=\s*(?:\(?\d{1,2}[.):]|[-•*–]\s|(?:comment|point|reviewer|referee)\s*#?\d+))", text, flags=re.I)
    points = []
    for part in parts:
        clean = re.sub(r"^\s*(?:\(?\d{1,2}[.):]|[-•*–]|(?:comment|point)\s*#?\d+\s*[:.)-]?)\s*", "", " ".join(part.split()), flags=re.I)
        if len(clean) >= 25:
            points.append({"comment": clean[:900], "segment": guess_segment(clean, keys), "instruction": clean[:600]})
    return points[:40]


def split_with_model(p: Project, comments: str, keys: list[str]) -> list[dict]:
    listing = "\n".join(f"- {k}: {p.segments[k].title} — begins: {' '.join(p.segments[k].content.split()[:30])}…" for k in keys)
    words = comments.split()
    points: list[dict] = []
    for start in range(0, min(len(words), 4800), 1200):   # long comment sheets are read in parts
        data = llm.generate_json(
            "You are an experienced journal editor helping an author plan a revision. Reply with JSON only.",
            f"ARTICLE: {p.title}\n\nSECTIONS OF THE ARTICLE (key: title):\n{listing}\n\n"
            f"COMMENTS RECEIVED FROM A REVIEWER OR SUPERVISOR:\n{' '.join(words[start:start + 1200])}\n\n"
            "Split the comments into separate points, one request each. Skip praise and remarks that ask for no change. For each point give:\n"
            "- comment: the reviewer's request in their own words, shortened to at most 60 words\n"
            "- segment: the ONE section key from the list that has to change to answer it\n"
            "- instruction: a precise instruction for revising that section (at most 50 words). Say what to add, remove or change. "
            "Never instruct to invent data, results or sources.\n"
            "JSON: {\"points\": [{\"comment\": \"\", \"segment\": \"\", \"instruction\": \"\"}]}",
            default={}, max_tokens=1500,
        )
        for pt in data.get("points") or [] if isinstance(data, dict) else []:
            if isinstance(pt, dict) and str(pt.get("comment", "")).strip() and str(pt.get("instruction", "")).strip():
                seg = pt.get("segment") if pt.get("segment") in keys else guess_segment(str(pt["comment"]), keys)
                points.append({"comment": str(pt["comment"]).strip()[:900], "segment": seg, "instruction": str(pt["instruction"]).strip()[:600]})
    return points[:40]


def analyse(pid: str, rid: str) -> None:
    """Job: fill the round's points. Falls back to a plain split when the model is unavailable or returns nothing."""
    rnd = get(pid, rid)
    p = store.load(pid)
    keys = editable(p)
    points, by = [], "model"
    try:
        if llm.health()["model_available"]:
            points = split_with_model(p, rnd["comments"], keys)
    except Exception:  # noqa: BLE001 - the plain split below still gives the author something to work with
        points = []
    if not points:
        points, by = split_plain(rnd["comments"], keys), "plain"
    rnd["points"] = [{"n": i, **pt, "decision": "accept", "state": "", "response": "", "change": ""} for i, pt in enumerate(points, 1)]
    rnd["status"], rnd["analysed_by"] = "ready", by
    save(pid, rnd)
    history.record(pid, "round.analysed", f"{rnd['name']}: comments split into {len(points)} point(s)", actor="RAF", round=rid)


# ---------------------------------------------------------------------------- 2. apply
def apply(pid: str, rid: str) -> None:
    """Job: apply every accepted point, one revision per segment, then draft the response letter."""
    rnd = get(pid, rid)
    p = store.load(pid)
    rnd["status"] = "applying"
    save(pid, rnd)
    before = revisions.snapshot(pid, f"{rnd['name']} — before", auto=True, trigger="round")
    rnd["before"] = before["n"] if before else None
    by_segment: dict[str, list[dict]] = {}
    for pt in rnd["points"]:
        if pt["decision"] == "accept" and pt["segment"] in p.segments and p.segments[pt["segment"]].content:
            by_segment.setdefault(pt["segment"], []).append(pt)
        else:
            pt["state"] = "skipped"

    def work(item):
        key, pts = item
        instruction = "Revise this section to answer the following reviewer comments. Address every one of them:\n" + "\n".join(
            f"{i}. {pt['instruction']}  (the reviewer wrote: “{pt['comment'][:280]}”)" for i, pt in enumerate(pts, 1))
        try:
            writer.revise_segment(pid, key, instruction, reason=f"revision round: {rnd['name']}")
            return key, ""
        except Exception as exc:  # noqa: BLE001 - one failing segment must not stop the round
            return key, str(exc)[:200]

    agents = max(1, min(3, int(p.options.get("agents", 1))))
    llm.runtime.ensure(agents)
    with ThreadPoolExecutor(max_workers=agents) as pool:
        for key, error in pool.map(work, by_segment.items()):
            for pt in by_segment[key]:
                pt["state"] = "failed" if error else "applied"
                if error:
                    pt["response"] = ""
                    rnd["error"] = f"{p.segments[key].title}: {error}"
    writer.sync_references(pid)
    after = revisions.snapshot(pid, f"{rnd['name']} — revised", trigger="round", actor="RAF")
    rnd["after"] = after["n"] if after else None
    describe_changes(pid, rnd)
    rnd["letter"] = letter(store.load(pid), rnd)
    rnd["status"], rnd["finished"] = "done", time.time()
    save(pid, rnd)
    done = sum(pt["state"] == "applied" for pt in rnd["points"])
    history.record(pid, "round.completed", f"{rnd['name']} completed — {done} of {len(rnd['points'])} point(s) applied in {len(by_segment)} segment(s)",
                   actor="RAF", round=rid, before=rnd["before"], after=rnd["after"])


def describe_changes(pid: str, rnd: dict) -> None:
    """Fill each applied point's `change` (what was added to its segment) and `response` (the reply to the reviewer)."""
    if not (rnd["before"] and rnd["after"]):
        return
    cmp = {s["key"]: s for s in revisions.compare(pid, rnd["before"], rnd["after"], rendered=True)["segments"]}
    p = store.load(pid)
    applied = [pt for pt in rnd["points"] if pt["state"] == "applied"]
    def stems(text: str) -> set[str]:
        return {w[:6] for w in re.findall(r"[a-z]{4,}", text.lower())}

    for pt in applied:
        seg = cmp.get(pt["segment"], {})
        # Quote the added sentences that speak to this comment, not simply the longest additions: several points can
        # share one segment, and the model may reword more than it was asked to.
        added = [s for op, t in seg.get("diff", []) if op == "+" for s in split_sentences(t.replace("\n\n", " ")) if len(s.split()) >= 6]
        asked = stems(pt["comment"] + " " + pt["instruction"])
        best = sorted(added, key=lambda s: (len(stems(s) & asked), len(s)), reverse=True)[:2]
        pt["change"] = " … ".join(" ".join(s.split()[:45]) for s in sorted(best, key=added.index))
        pt["words"] = [seg.get("added", 0), seg.get("removed", 0)]
    for i in range(0, len(applied), 5):
        batch = applied[i:i + 5]
        blocks = "\n\n".join(
            f"POINT {pt['n']}\nReviewer comment: {pt['comment']}\nSection revised: {p.segments[pt['segment']].title}\n"
            f"Text added in the revision: {pt['change'] or '(the section was reworded; no long passage was added)'}" for pt in batch)
        data = llm.generate_json(
            "You write the author's replies in a response-to-reviewers letter. Reply with JSON only.",
            f"ARTICLE: {p.title}\n\n{blocks}\n\nFor each point write the author's reply: courteous, specific, 2–3 sentences. Say what was "
            "changed and in which section, using ONLY the added text shown. Do not claim changes that are not shown and do not quote more than a few words.\n"
            "JSON: {\"responses\": [{\"point\": 1, \"response\": \"\"}]}",
            default={}, max_tokens=220 * len(batch) + 150,
        )
        for r in data.get("responses") or [] if isinstance(data, dict) else []:
            for pt in batch:
                if isinstance(r, dict) and str(r.get("point")) == str(pt["n"]) and str(r.get("response", "")).strip():
                    pt["response"] = str(r["response"]).strip()[:900]
    for pt in applied:
        pt["response"] = pt["response"] or f"We thank the reviewer for this comment. We have revised the {p.segments[pt['segment']].title} to address it."


# ---------------------------------------------------------------------------- 3. letter
def letter(p: Project, rnd: dict) -> str:
    """The response letter as Markdown. The author edits this text; the DOCX is made from whatever it says."""
    title = p.segments["title"].content.strip() if p.segments.get("title") and p.segments["title"].content else p.title
    lines = ["# Response to reviewers", "", f"**{title}**", "", f"{rnd['name']} · {time.strftime('%d %B %Y')}", "",
             "We thank the reviewers for their careful reading of the manuscript. Each comment is reproduced below, followed by our response "
             "and the change made. A version of the manuscript with all changes highlighted accompanies this letter.", ""]
    for pt in rnd["points"]:
        seg = p.segments.get(pt["segment"])
        lines += [f"## Comment {pt['n']}", "", f"> {pt['comment']}", ""]
        if pt["state"] == "applied":
            lines += [f"**Response.** {pt['response']}", ""]
            if pt.get("change"):
                lines += [f"**Change made** ({seg.title if seg else pt['segment']}). “{pt['change']}”", ""]
        elif pt["state"] == "failed":
            lines += ["**Response.** [This change could not be applied automatically — revise the section in the Studio and describe the change here.]", ""]
        else:
            lines += ["**Response.** [No change was made for this comment — explain why here.]", ""]
    return "\n".join(lines)


def letter_docx(text: str, out: Path) -> Path:
    import docx
    from docx.shared import Cm, Pt

    d = docx.Document()
    d.styles["Normal"].font.name, d.styles["Normal"].font.size = "Cambria", Pt(11)
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("## "):
            d.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            d.add_heading(line[2:], level=0)
        else:
            par = d.add_paragraph()
            quote = line.startswith(">")
            if quote:
                line = line.lstrip("> ")
                par.paragraph_format.left_indent = Cm(1)
            for i, piece in enumerate(re.split(r"\*\*(.+?)\*\*", line)):
                if piece:
                    r = par.add_run(piece)
                    r.bold, r.italic = i % 2 == 1, quote
    d.save(out)
    return out
