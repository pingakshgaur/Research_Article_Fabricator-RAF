"""Article Revisions: saved states of the whole article, with compare, restore and a changes-highlighted export.

Stored in  data/projects/<id>/revisions/r0001.json  (+ r0001_figures/ with the charts as they were).

A revision holds every segment as it stood (text, status, tables, figures, quality report) but not the per-segment
version lists. RAF saves one automatically before anything that replaces a lot of text at once — a new fabrication run,
an import, a restore, a revision round — and at every publish. Restoring never destroys work: the current state is
saved as a revision first.
"""
from __future__ import annotations

import difflib
import json
import re
import shutil
import time
from pathlib import Path

from . import citations, history, store, structure
from .checkpoints import signature
from .models import CustomSegment, Project, Segment, Version

KEEP_AUTO = 30
_PARA = "\n\n"


def _dir(pid: str) -> Path:
    d = store.project_dir(pid) / "revisions"
    d.mkdir(exist_ok=True)
    return d


def _order(p: Project, keys) -> list[str]:
    keys = list(keys)
    known = structure.order(p, keys)
    return known + [k for k in keys if k not in known]      # segments of an old revision that the project no longer defines go last


def _state(p: Project) -> dict:
    """The article as it is now, in the same shape as a saved revision."""
    segs = {k: s.model_dump(exclude={"versions"}) for k, s in p.segments.items() if s.content}
    return {
        "title": (p.segments["title"].content.strip() if p.segments.get("title") and p.segments["title"].content else p.title),
        "citation_style": p.citation_style, "selected": list(p.selected), "segments": segs,
        "structure": {"titles": dict(p.titles), "custom_segments": {k: c.model_dump() for k, c in p.custom_segments.items()},
                      "order": structure.order(p), "excluded": list(p.excluded)},
        "signature": signature({k: s["content"] for k, s in sorted(segs.items())}),
        "stats": {"segments": len(segs), "approved": sum(s["status"] == "approved" for s in segs.values()),
                  "words": sum(s["quality"]["words"] for k, s in segs.items() if k in store.BODY),
                  "citations": sum(s["quality"]["citations"] for s in segs.values())},
    }


def _meta(rev: dict) -> dict:
    return {k: v for k, v in rev.items() if k != "segments"}


def list_(pid: str) -> list[dict]:
    """Saved revisions without their text, newest first."""
    out = []
    for f in sorted(_dir(pid).glob("r[0-9]*.json"), reverse=True):
        try:
            out.append(_meta(json.loads(f.read_text(encoding="utf-8"))))
        except json.JSONDecodeError:
            continue
    return out


def get(pid: str, n: int) -> dict:
    path = _dir(pid) / f"r{n:04d}.json"
    if not path.exists():
        raise KeyError(n)
    return json.loads(path.read_text(encoding="utf-8"))


def snapshot(pid: str, name: str, *, auto: bool = False, trigger: str = "manual", note: str = "", actor: str = "") -> dict | None:
    """Save the article as it is now. Returns the revision (without text), or None when there is nothing written yet.

    An automatic snapshot of a state that is already the latest revision is not saved twice; that revision is returned.
    """
    with store.lock_for(pid):
        p = store.load(pid)
        if p is None:
            raise KeyError(pid)
        state = _state(p)
        if not state["segments"]:
            return None
        existing = list_(pid)
        if auto and existing and existing[0]["signature"] == state["signature"]:
            return existing[0]
        n = (existing[0]["n"] + 1) if existing else 1
        rev = {"n": n, "name": name.strip()[:120] or f"Revision {n}", "note": note.strip()[:500], "auto": auto, "trigger": trigger,
               "created": time.time(), **state}
        d = _dir(pid)
        figures = [f for f in (store.project_dir(pid) / "figures").glob("*") if f.is_file()]
        if figures:
            (d / f"r{n:04d}_figures").mkdir(exist_ok=True)
            for f in figures:
                shutil.copy2(f, d / f"r{n:04d}_figures" / f.name)
        (d / f"r{n:04d}.json").write_text(json.dumps(rev, ensure_ascii=False), encoding="utf-8")
        for old in [r for r in existing if r["auto"]][KEEP_AUTO - 1:] if auto else []:
            _remove(pid, old["n"])
    history.record(pid, "revision.saved", f"Revision saved — “{rev['name']}”" + (" (automatic)" if auto else ""),
                   actor=actor or ("RAF" if auto else "you"), revision=n, words=state["stats"]["words"], trigger=trigger)
    return _meta(rev)


def _remove(pid: str, n: int) -> None:
    d = _dir(pid)
    (d / f"r{n:04d}.json").unlink(missing_ok=True)
    shutil.rmtree(d / f"r{n:04d}_figures", ignore_errors=True)


def delete(pid: str, n: int) -> None:
    rev = get(pid, n)
    _remove(pid, n)
    history.record(pid, "revision.deleted", f"Revision deleted — “{rev['name']}”", revision=n)


def rename(pid: str, n: int, name: str, note: str | None = None) -> dict:
    rev = get(pid, n)
    rev["name"] = name.strip()[:120] or rev["name"]
    if note is not None:
        rev["note"] = note.strip()[:500]
    rev["auto"] = False   # a revision the author named is kept like one they saved
    (_dir(pid) / f"r{n:04d}.json").write_text(json.dumps(rev, ensure_ascii=False), encoding="utf-8")
    return _meta(rev)


def restore(pid: str, n: int, key: str | None = None) -> Project:
    """Bring back a whole revision, or one segment of it. The current state is saved as a revision first."""
    rev = get(pid, n)
    if key and key not in rev["segments"]:
        raise KeyError(key)
    snapshot(pid, f"Before restoring “{rev['name']}”", auto=True, trigger="restore")
    saved_figures = _dir(pid) / f"r{n:04d}_figures"
    keys = [key] if key else list(rev["segments"])
    if saved_figures.exists() and any(rev["segments"][k].get("figures") for k in keys):
        for f in saved_figures.glob("*"):
            shutil.copy2(f, store.project_dir(pid) / "figures" / f.name)
    reason = f"restored from revision “{rev['name']}”"[:120]

    saved = rev.get("structure") or {}     # revisions made before 1.6 carry no structure

    def apply(p: Project):
        for k in keys:   # a custom segment that was deleted since the revision gets its definition back
            if k in saved.get("custom_segments", {}) and (not key or k not in p.custom_segments):
                p.custom_segments[k] = CustomSegment(**saved["custom_segments"][k])
        if not key and saved:
            p.titles, p.order, p.excluded = dict(saved["titles"]), list(saved["order"]), list(saved["excluded"])
            p.custom_segments = {k: CustomSegment(**c) for k, c in saved["custom_segments"].items()}
        for k in keys:
            data = rev["segments"][k]
            old = p.segments.get(k)
            seg = Segment(**data, versions=[*(old.versions if old else []), Version(content=data["content"], reason=reason)][-20:])
            seg.status, seg.error = ("approved" if data["status"] == "approved" else "draft"), ""
            p.segments[k] = seg
        if not key:   # the article becomes the revision: segments written since then leave (they are in the snapshot above)
            for k in [k for k in p.segments if k not in rev["segments"]]:
                del p.segments[k]
            p.selected = [k for k in rev.get("selected", []) if k in p.segments] or _order(p, p.segments)
            if p.stage in {"setup", "references", "segments", "processing"}:
                p.stage = "studio"
    p = store.mutate(pid, apply)
    what = f"{rev['segments'][key]['title']} restored" if key else "Whole article restored"
    history.record(pid, "revision.restored", f"{what} from revision “{rev['name']}”", segment=key or "", revision=n)
    return p


# ---------------------------------------------------------------------------- compare
def _tokens(text: str) -> list[str]:
    return re.findall(r"\n\s*\n|\S+", text)


def _join(tokens: list[str]) -> str:
    out = ""
    for t in tokens:
        out += _PARA if t.strip() == "" else (t if not out or out.endswith(_PARA) else " " + t)
    return out


def word_diff(a: str, b: str) -> list[list[str]]:
    """[["=", text], ["-", removed], ["+", added], …] at word level; paragraph breaks are kept as "\\n\\n"."""
    ta, tb = _tokens(a), _tokens(b)
    out: list[list[str]] = []
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, ta, tb, autojunk=False).get_opcodes():
        if op == "equal":
            out.append(["=", _join(ta[i1:i2])])
            continue
        if i2 > i1:
            out.append(["-", _join(ta[i1:i2])])
        if j2 > j1:
            out.append(["+", _join(tb[j1:j2])])
    return out


def _side(pid: str, which: str | int, p: Project) -> dict:
    if str(which) == "current":
        return {"n": "current", "name": "Current article", "created": time.time(), "auto": False, **_state(p)}
    return get(pid, int(which))


def compare(pid: str, a: str | int, b: str | int, *, rendered: bool = False) -> dict:
    """Differences between two revisions (or a revision and the current article), segment by segment.

    `rendered=True` shows citations as they are published ("(Smith, 2021)") instead of RAF's markers ("[R3]").
    """
    p = store.load(pid)
    if p is None:
        raise KeyError(pid)
    ra, rb = _side(pid, a, p), _side(pid, b, p)

    def text(rev: dict, key: str) -> str:
        t = rev["segments"].get(key, {}).get("content", "")
        t = re.sub(r"\[\[(TABLE|FIGURE) (\d+)\]\]", lambda m: f"[{m.group(1).capitalize()} {m.group(2)}]", t)
        if rendered and key != "references":
            order = citations.reference_list(p, [s["content"] for k, s in rev["segments"].items() if k != "references"])[1]
            t = citations.render_in_text(t, p, order)
        return t

    segments, added, removed = [], 0, 0
    for key in _order(p, set(ra["segments"]) | set(rb["segments"])):
        ta, tb = text(ra, key), text(rb, key)
        status = "added" if not ta else "removed" if not tb else "same" if ta == tb else "changed"
        diff = word_diff(ta, tb) if status != "same" else [["=", tb]]
        plus = sum(len(t.split()) for op, t in diff if op == "+")
        minus = sum(len(t.split()) for op, t in diff if op == "-")
        added, removed = added + plus, removed + minus
        seg = (rb["segments"].get(key) or ra["segments"][key])
        segments.append({"key": key, "title": seg["title"], "status": status, "added": plus, "removed": minus, "diff": diff,
                         "in_a": bool(ta), "in_b": bool(tb)})
    return {"a": _meta(ra), "b": _meta(rb), "segments": segments,
            "totals": {"added": added, "removed": removed, "changed_segments": sum(s["status"] != "same" for s in segments)}}


# ---------------------------------------------------------------------------- changes-highlighted export
INS, DEL = "0A7A3D", "B00020"


def export_changes(pid: str, a: str | int, b: str | int, fmt: str, out: Path) -> Path:
    """A DOCX or PDF of the differences: insertions green and underlined, deletions red and struck through."""
    cmp = compare(pid, a, b, rendered=True)
    title = cmp["b"]["title"]
    subtitle = f"Changes from “{cmp['a']['name']}” to “{cmp['b']['name']}”"
    totals = f"{cmp['totals']['added']:,} words added · {cmp['totals']['removed']:,} words removed · {cmp['totals']['changed_segments']} segment(s) changed"
    shown = [s for s in cmp["segments"] if s["status"] != "same"]

    def paragraphs(diff):   # split the op list at paragraph breaks -> [[(op, text), …], …]
        paras, cur = [], []
        for op, text in diff:
            pieces = text.split(_PARA)
            for i, piece in enumerate(pieces):
                if i:
                    paras.append(cur)
                    cur = []
                if piece.strip():
                    cur.append((op, piece.strip()))
        paras.append(cur)
        return [x for x in paras if x]

    if fmt == "docx":
        import docx
        from docx.shared import Pt, RGBColor

        d = docx.Document()
        d.styles["Normal"].font.name, d.styles["Normal"].font.size = "Cambria", Pt(10.5)
        d.add_heading(title, level=0)
        d.add_paragraph(subtitle).runs[0].italic = True
        legend = d.add_paragraph(totals + " · ")
        r = legend.add_run("added text")
        r.font.color.rgb, r.font.underline = RGBColor.from_string(INS), True
        legend.add_run(" · ")
        r = legend.add_run("removed text")
        r.font.color.rgb, r.font.strike = RGBColor.from_string(DEL), True
        for s in shown:
            d.add_heading(f"{s['title']}  (+{s['added']} / −{s['removed']} words)", level=1)
            for para in paragraphs(s["diff"]):
                par = d.add_paragraph()
                for i, (op, text) in enumerate(para):
                    r = par.add_run(("" if i == 0 else " ") + text)
                    if op == "+":
                        r.font.color.rgb, r.font.underline = RGBColor.from_string(INS), True
                    elif op == "-":
                        r.font.color.rgb, r.font.strike = RGBColor.from_string(DEL), True
        if not shown:
            d.add_paragraph("The two versions are identical.")
        d.save(out)
        return out

    from xml.sax.saxutils import escape

    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    from .publisher import pdf_fonts

    fonts = pdf_fonts()
    body = ParagraphStyle("b", fontName=fonts["serif"][0], fontSize=10.5, leading=15, spaceAfter=7)
    h1 = ParagraphStyle("h", fontName=fonts["sans"][1], fontSize=13, leading=17, spaceBefore=12, spaceAfter=6)
    head = ParagraphStyle("t", fontName=fonts["sans"][1], fontSize=18, leading=22, spaceAfter=4)
    mark = {"+": f"<font color='#{INS}'><u>%s</u></font>", "-": f"<font color='#{DEL}'><strike>%s</strike></font>", "=": "%s"}
    story = [Paragraph(escape(title), head), Paragraph(f"<i>{escape(subtitle)}</i>", body),
             Paragraph(escape(totals) + " · " + mark["+"] % "added text" + " · " + mark["-"] % "removed text", body), Spacer(1, 6)]
    for s in shown:
        story.append(Paragraph(escape(f"{s['title']}  (+{s['added']} / −{s['removed']} words)"), h1))
        story += [Paragraph(" ".join(mark[op] % escape(text) for op, text in para), body) for para in paragraphs(s["diff"])]
    if not shown:
        story.append(Paragraph("The two versions are identical.", body))
    SimpleDocTemplate(str(out), pagesize=A4, leftMargin=2.3 * cm, rightMargin=2.3 * cm, topMargin=2.2 * cm, bottomMargin=2.2 * cm,
                      title=f"{title} — changes").build(story)
    return out
