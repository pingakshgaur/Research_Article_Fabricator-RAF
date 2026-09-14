"""Publishing: assemble approved segments into a professional DOCX, PDF or Markdown document using a chosen template and
colour style, and re-import edited DOCX / Markdown / PDF files back into segments."""
from __future__ import annotations

import base64
import datetime as dt
import re
from dataclasses import dataclass, field
from pathlib import Path

from . import citations, store
from .blueprints import BLUEPRINTS
from .models import Figure, Project, Table

PLACEHOLDER = re.compile(r"^\[\[(TABLE|FIGURE) (\d+)\]\]$")


# ====================================================================== templates & palettes
@dataclass(frozen=True)
class Template:
    key: str
    name: str
    description: str
    body_serif: bool = True
    heading_serif: bool = False
    body_pt: float = 10.5
    leading: float = 1.45
    margins_cm: tuple = (2.3, 2.3, 2.2, 2.3)           # top, right, bottom, left
    columns: int = 1
    title_align: str = "left"                            # left | center
    cover: str = "band"                                  # band | rule | plain | title_page
    numbering: str = "arabic"                            # arabic | roman | none
    heading_upper: bool = False
    heading_center: bool = False
    heading_scale: float = 3.5                           # points added to body size for level-1 headings
    abstract_style: str = "box"                          # box | rules | plain | ieee | page
    table_style: str = "shaded"                          # shaded | booktabs | grid
    justify: bool = True
    indent_cm: float = 0.0
    toc: bool = True
    color: bool = True
    kicker: str = "RESEARCH ARTICLE"
    running_header: bool = True
    docx_body_font: str = "Cambria"
    docx_heading_font: str = "Calibri"


TEMPLATES: dict[str, Template] = {
    "modern_report": Template(
        "modern_report", "Modern Report", "Colour band, sans-serif headings, shaded abstract and tables, table of contents."),
    "classic_journal": Template(
        "classic_journal", "Classic Journal", "Centred serif title, ruled abstract, book-style tables and indented paragraphs.",
        heading_serif=True, body_pt=10.5, leading=1.4, margins_cm=(2.5, 2.4, 2.4, 2.4), title_align="center", cover="rule",
        abstract_style="rules", table_style="booktabs", indent_cm=0.6, toc=False, kicker="", heading_scale=2.5,
        docx_body_font="Georgia", docx_heading_font="Georgia"),
    "apa_manuscript": Template(
        "apa_manuscript", "APA 7 Manuscript", "Submission-ready: title page, double spacing, 12 pt serif, centred headings, monochrome.",
        heading_serif=True, body_pt=12, leading=2.0, margins_cm=(2.54, 2.54, 2.54, 2.54), title_align="center", cover="title_page",
        numbering="none", heading_center=True, heading_scale=0, abstract_style="page", table_style="booktabs", justify=False,
        indent_cm=1.27, toc=False, color=False, kicker="", docx_body_font="Times New Roman", docx_heading_font="Times New Roman"),
    "conference": Template(
        "conference", "Two-Column Conference", "IEEE-style two-column body, run-in abstract, Roman-numeral small-caps headings.",
        heading_serif=True, body_pt=9.5, leading=1.25, margins_cm=(1.9, 1.6, 2.2, 1.6), columns=2, title_align="center",
        cover="plain", numbering="roman", heading_upper=True, heading_center=True, heading_scale=0.5, abstract_style="ieee",
        table_style="grid", indent_cm=0.35, toc=False, kicker="", running_header=False,
        docx_body_font="Times New Roman", docx_heading_font="Times New Roman"),
    "minimal_monograph": Template(
        "minimal_monograph", "Minimal Monograph", "Generous margins, large serif headings, ragged-right text and quiet colour accents.",
        heading_serif=True, body_pt=11, leading=1.6, margins_cm=(3.0, 3.2, 3.0, 3.2), cover="plain", abstract_style="plain",
        table_style="booktabs", justify=False, heading_scale=7, kicker="", docx_body_font="Georgia", docx_heading_font="Georgia"),
}

PALETTES: dict[str, dict] = {
    "violet": {"name": "Violet", "accent": "3B1782", "fill": "F3EFFB", "rule": "C9BFE0"},
    "midnight": {"name": "Midnight Blue", "accent": "1E3A8A", "fill": "EEF2FF", "rule": "C7D2FE"},
    "emerald": {"name": "Emerald", "accent": "065F46", "fill": "ECFDF5", "rule": "A7F3D0"},
    "crimson": {"name": "Crimson", "accent": "9F1239", "fill": "FFF1F2", "rule": "FECDD3"},
    "ochre": {"name": "Ochre", "accent": "92400E", "fill": "FFFBEB", "rule": "FDE68A"},
    "lime": {"name": "Forest Lime", "accent": "3F6212", "fill": "F7FEE7", "rule": "D9F99D"},
    "graphite": {"name": "Graphite", "accent": "1F2937", "fill": "F3F4F6", "rule": "D1D5DB"},
}
MONO = {"name": "Monochrome", "accent": "000000", "fill": "FFFFFF", "rule": "000000"}


def styles_catalogue() -> dict:
    return {
        "templates": [{"key": t.key, "name": t.name, "description": t.description, "columns": t.columns, "color": t.color,
                       "cover": t.cover, "serif": t.heading_serif} for t in TEMPLATES.values()],
        "palettes": [{"key": k, **v} for k, v in PALETTES.items()],
    }


def resolve_style(template: str | None, palette: str | None) -> tuple[Template, dict]:
    tpl = TEMPLATES.get(template or "", TEMPLATES["modern_report"])
    pal = PALETTES.get(palette or "", PALETTES["violet"]) if tpl.color else MONO
    return tpl, pal


ROMAN = ["", "I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]


def heading_text(tpl: Template, number: str, heading: str) -> str:
    label = heading.upper() if tpl.heading_upper else heading
    if not number or tpl.numbering == "none":
        return label
    if tpl.numbering == "roman":
        return f"{ROMAN[int(number)] if int(number) < len(ROMAN) else number}. {label}"
    return f"{number}. {label}"


# ====================================================================== article model
@dataclass
class Block:
    kind: str                  # "p" | "table" | "figure" | "ref" | "subhead"
    text: str = ""
    table: Table | None = None
    figure: Figure | None = None


@dataclass
class Section:
    key: str
    heading: str
    number: str
    blocks: list[Block] = field(default_factory=list)


@dataclass
class Article:
    title: str
    subtitle: str
    authors: list[str]
    affiliation: str
    date: str
    abstract: str
    keywords: str
    sections: list[Section]
    project_dir: Path


def assemble(p: Project, only_approved: bool = False) -> Article:
    segs = {k: s for k, s in p.segments.items() if s.content and (s.status == "approved" or not only_approved)}
    body_texts = [s.content for k, s in segs.items() if k != "references"]
    _, order = citations.reference_list(p, body_texts)

    title_text = segs["title"].content.strip() if "title" in segs else p.title
    title, _, subtitle = title_text.partition(": ")
    abstract = citations.render_in_text(segs["abstract"].content, p, order) if "abstract" in segs else ""
    keywords = segs["keywords"].content if "keywords" in segs else ""

    sections: list[Section] = []
    numbered = 0
    for key in sorted(segs, key=lambda k: BLUEPRINTS[k].display_order):
        if key in {"title", "abstract", "keywords"}:
            continue
        seg = segs[key]
        if key in {"references", "appendices"}:
            number = ""
        else:
            numbered += 1
            number = str(numbered)
        sec = Section(key=key, heading=seg.title, number=number)
        tables = {t.number: t for t in seg.tables}
        figures = {f.number: f for f in seg.figures}
        for raw in re.split(r"\n\s*\n", seg.content):
            raw = raw.strip()
            if not raw:
                continue
            m = PLACEHOLDER.match(raw)
            if m:
                kind, num = m.group(1), int(m.group(2))
                if kind == "TABLE" and num in tables:
                    sec.blocks.append(Block("table", table=tables[num]))
                elif kind == "FIGURE" and num in figures:
                    sec.blocks.append(Block("figure", figure=figures[num]))
                continue
            if key == "references":
                for line in raw.splitlines():
                    if line.strip():
                        sec.blocks.append(Block("ref", text=line.strip()))
                continue
            if key == "appendices" and re.match(r"^Appendix [A-Z]\.", raw):
                sec.blocks.append(Block("subhead", text=raw))
                continue
            for para in raw.split("\n") if key == "appendices" else [raw]:
                text = citations.render_in_text(para.strip().lstrip("- "), p, order)
                if text:
                    sec.blocks.append(Block("p", text=text))
        sections.append(sec)
    return Article(title=title.strip(), subtitle=subtitle.strip(), authors=p.authors, affiliation=p.affiliation,
                   date=dt.date.today().strftime("%d %B %Y"), abstract=abstract, keywords=keywords, sections=sections,
                   project_dir=store.project_dir(p.id))


def table_label(t: Table) -> str:
    return f"Table A{t.number - 100}" if t.number > 100 else f"Table {t.number}"


def _clean_caption(caption: str) -> str:
    return caption.rstrip(".") + "."


# ====================================================================== DOCX
def to_docx(a: Article, out: Path, template: str | None = None, palette: str | None = None) -> Path:
    import docx
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    tpl, pal = resolve_style(template, palette)
    accent = RGBColor.from_string(pal["accent"])
    muted = RGBColor(0x55, 0x55, 0x60) if tpl.color else RGBColor(0, 0, 0)
    body_align = WD_ALIGN_PARAGRAPH.JUSTIFY if tpl.justify else WD_ALIGN_PARAGRAPH.LEFT
    title_align = WD_ALIGN_PARAGRAPH.CENTER if tpl.title_align == "center" else WD_ALIGN_PARAGRAPH.LEFT

    d = docx.Document()
    sec = d.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin, sec.right_margin, sec.bottom_margin, sec.left_margin = (Cm(v) for v in tpl.margins_cm)
    col_width_cm = (21.0 - tpl.margins_cm[1] - tpl.margins_cm[3] - (0.75 if tpl.columns == 2 else 0)) / tpl.columns

    styles = d.styles
    normal = styles["Normal"]
    normal.font.name, normal.font.size = tpl.docx_body_font, Pt(tpl.body_pt)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), tpl.docx_body_font)
    normal.paragraph_format.space_after = Pt(0 if tpl.leading >= 2 else 6)
    normal.paragraph_format.line_spacing = tpl.leading
    for lvl in (1, 2):
        h = styles[f"Heading {lvl}"]
        size = tpl.body_pt + (tpl.heading_scale if lvl == 1 else max(0.5, tpl.heading_scale / 2))
        h.font.name, h.font.size, h.font.bold = tpl.docx_heading_font, Pt(size), True
        h.font.italic = False
        h.font.color.rgb = accent if tpl.color else RGBColor(0, 0, 0)
        h.font.all_caps = tpl.heading_upper and lvl == 1
        h.element.rPr.rFonts.set(qn("w:eastAsia"), tpl.docx_heading_font)
        h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(16 if lvl == 1 else 10), Pt(6)
        h.paragraph_format.keep_with_next = True
        h.paragraph_format.line_spacing = tpl.leading if tpl.leading >= 2 else 1.1
        if tpl.heading_center and lvl == 1:
            h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def field(paragraph, instr: str):
        run = paragraph.add_run()
        for tag, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
            elif text == instr:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = instr
            else:
                el = OxmlElement("w:t")
                el.text = text
            run._r.append(el)
        return run

    def shade(cell, hex_fill: str):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_fill)
        tcPr.append(shd)

    def para_rule(paragraph, edge="bottom", color=None, size="8"):
        pPr = paragraph._p.get_or_add_pPr()
        bdr = pPr.find(qn("w:pBdr"))
        if bdr is None:
            bdr = OxmlElement("w:pBdr")
            pPr.append(bdr)
        b = OxmlElement(f"w:{edge}")
        for k, v in (("w:val", "single"), ("w:sz", size), ("w:space", "4"), ("w:color", color or pal["rule"])):
            b.set(qn(k), v)
        bdr.append(b)

    def table_borders(table):
        tblPr = table._tbl.tblPr
        el = OxmlElement("w:tblBorders")
        if tpl.table_style == "grid":
            edges = {e: ("single", "4", "888888") for e in ("top", "left", "bottom", "right", "insideH", "insideV")}
        elif tpl.table_style == "booktabs":
            edges = {"top": ("single", "12", pal["accent"] if tpl.color else "000000"), "bottom": ("single", "12", pal["accent"] if tpl.color else "000000")}
        else:
            edges = {"top": ("single", "12", pal["accent"]), "bottom": ("single", "12", pal["accent"]), "insideH": ("single", "4", pal["rule"])}
        for edge, (val, sz, color) in edges.items():
            e = OxmlElement(f"w:{edge}")
            for k, v in (("w:val", val), ("w:sz", sz), ("w:color", color)):
                e.set(qn(k), v)
            el.append(e)
        tblPr.append(el)

    def cell_bottom(cell, color, size="6"):
        tcPr = cell._tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        e = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", size), ("w:color", color)):
            e.set(qn(k), v)
        b.append(e)
        tcPr.append(b)

    def set_columns(section, num: int):
        cols = section._sectPr.find(qn("w:cols"))
        if cols is None:
            cols = OxmlElement("w:cols")
            section._sectPr.append(cols)
        cols.set(qn("w:num"), str(num))
        cols.set(qn("w:space"), "425")

    def run(par, text, size=None, bold=None, italic=None, color=None, font=None):
        r = par.add_run(text)
        if size:
            r.font.size = Pt(size)
        if bold is not None:
            r.font.bold = bold
        if italic is not None:
            r.font.italic = italic
        if color is not None:
            r.font.color.rgb = color
        if font:
            r.font.name = font
        return r

    # ---- header / footer
    if tpl.cover == "title_page":      # APA: page number top right
        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        field(hp, "PAGE")
    else:
        if tpl.running_header:
            hp = sec.header.paragraphs[0]
            hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run(hp, a.title if len(a.title) < 90 else a.title[:87] + "…", size=8.5, color=muted, font=tpl.docx_heading_font)
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        field(fp, "PAGE").font.size = Pt(9)

    # ---- title block
    title_font = tpl.docx_heading_font
    if tpl.cover == "title_page":
        for _ in range(6):
            d.add_paragraph()
        t = d.add_paragraph()
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(t, a.title + (f": {a.subtitle}" if a.subtitle else ""), size=tpl.body_pt, bold=True)
        d.add_paragraph()
        for line in [", ".join(a.authors), a.affiliation, a.date]:
            if line:
                m = d.add_paragraph()
                m.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run(m, line)
        d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    else:
        if tpl.kicker:
            k = d.add_paragraph()
            k.paragraph_format.space_before = Pt(30)
            k.alignment = title_align
            run(k, tpl.kicker, size=9, bold=True, color=accent, font=title_font)
        t = d.add_paragraph()
        t.alignment = title_align
        t.paragraph_format.space_after = Pt(4)
        t.paragraph_format.line_spacing = 1.05
        run(t, a.title, size=(18 if tpl.columns == 2 else 22), bold=True, font=title_font)
        if a.subtitle:
            s = d.add_paragraph()
            s.alignment = title_align
            run(s, a.subtitle, size=(12 if tpl.columns == 2 else 14), italic=True, color=muted)
        meta = d.add_paragraph()
        meta.alignment = title_align
        if a.authors:
            run(meta, ", ".join(a.authors), size=tpl.body_pt + 0.5, bold=True)
            meta.add_run("\n")
        if a.affiliation:
            run(meta, a.affiliation + "\n", size=tpl.body_pt - 0.5, italic=True)
        run(meta, a.date, size=tpl.body_pt - 1, color=muted)
        if tpl.cover in {"band", "rule"}:
            para_rule(meta, "bottom", pal["accent"], "12")

    # ---- abstract
    if a.abstract:
        if tpl.abstract_style == "page":
            h = d.add_paragraph()
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run(h, "Abstract", bold=True)
            ab = d.add_paragraph(a.abstract)
            ab.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if a.keywords:
                kw = d.add_paragraph()
                kw.paragraph_format.first_line_indent = Cm(1.27)
                run(kw, "Keywords: ", italic=True)
                run(kw, a.keywords.replace(";", ","))
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            t = d.add_paragraph()
            t.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run(t, a.title + (f": {a.subtitle}" if a.subtitle else ""), bold=True)
        elif tpl.abstract_style == "box":
            box = d.add_table(rows=1, cols=1)
            box.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = box.rows[0].cells[0]
            shade(cell, pal["fill"])
            run(cell.paragraphs[0], "Abstract", size=tpl.body_pt + 0.5, bold=True, color=accent, font=title_font)
            body = cell.add_paragraph(a.abstract)
            body.alignment = body_align
            if a.keywords:
                kw = cell.add_paragraph()
                run(kw, "Keywords: ", size=tpl.body_pt - 0.5, bold=True)
                run(kw, a.keywords, size=tpl.body_pt - 0.5, italic=True)
            d.add_paragraph()
        elif tpl.abstract_style == "ieee":
            ab = d.add_paragraph()
            ab.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run(ab, "Abstract—", bold=True, italic=True, size=tpl.body_pt)
            run(ab, a.abstract, bold=True, size=tpl.body_pt)
            if a.keywords:
                kw = d.add_paragraph()
                run(kw, "Index Terms—", bold=True, italic=True, size=tpl.body_pt)
                run(kw, a.keywords.replace(";", ","), bold=True, size=tpl.body_pt)
        else:  # rules / plain
            h = d.add_paragraph()
            h.alignment = title_align if tpl.abstract_style == "rules" else WD_ALIGN_PARAGRAPH.LEFT
            run(h, "Abstract" if tpl.abstract_style == "plain" else "ABSTRACT", size=tpl.body_pt + (2 if tpl.abstract_style == "plain" else 0),
                bold=True, color=accent, font=title_font)
            if tpl.abstract_style == "rules":
                para_rule(h, "top", pal["accent"], "8")
            ab = d.add_paragraph(a.abstract)
            ab.alignment = body_align
            if a.keywords:
                kw = d.add_paragraph()
                run(kw, "Keywords: ", bold=True, size=tpl.body_pt - 0.5)
                run(kw, a.keywords, italic=True, size=tpl.body_pt - 0.5)
                if tpl.abstract_style == "rules":
                    para_rule(kw, "bottom", pal["accent"], "8")

    if tpl.toc:
        toc_title = d.add_paragraph()
        run(toc_title, "Contents", size=tpl.body_pt + 1.5, bold=True, color=accent, font=title_font)
        field(d.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
        note = d.add_paragraph()
        run(note, "(Right-click → Update Field to refresh the table of contents.)", size=8, italic=True, color=muted)
        d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    if tpl.columns == 2:
        body_section = d.add_section(WD_SECTION.CONTINUOUS)
        set_columns(body_section, 2)

    # ---- sections
    for s in a.sections:
        if s.key == "appendices" and tpl.columns == 1:
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        d.add_heading(heading_text(tpl, s.number, s.heading), level=1)
        first = True
        for b in s.blocks:
            if b.kind == "p":
                para = d.add_paragraph(b.text)
                para.alignment = body_align
                if tpl.indent_cm and not (first and tpl.key == "classic_journal"):
                    para.paragraph_format.first_line_indent = Cm(tpl.indent_cm)
                first = False
            elif b.kind == "subhead":
                d.add_heading(b.text, level=2)
            elif b.kind == "ref":
                para = d.add_paragraph(b.text)
                para.paragraph_format.left_indent = Cm(1.27 if tpl.cover == "title_page" else 1.0)
                para.paragraph_format.first_line_indent = Cm(-1.27 if tpl.cover == "title_page" else -1.0)
                para.paragraph_format.space_after = Pt(0 if tpl.leading >= 2 else 4)
                for r in para.runs:
                    r.font.size = Pt(tpl.body_pt - (0 if tpl.cover == "title_page" else 1))
            elif b.kind == "table" and b.table:
                tb = b.table
                cap = d.add_paragraph()
                cap.paragraph_format.keep_with_next = True
                if tpl.cover == "title_page":   # APA: bold label line, italic title line
                    run(cap, f"{table_label(tb)}\n", bold=True)
                    run(cap, tb.caption.rstrip("."), italic=True)
                else:
                    run(cap, f"{table_label(tb)}. ", size=tpl.body_pt - 0.5, bold=True, color=accent)
                    run(cap, _clean_caption(tb.caption), size=tpl.body_pt - 0.5, italic=True)
                cols = max(len(tb.columns), max((len(r) for r in tb.rows), default=0))
                table = d.add_table(rows=1, cols=cols)
                if tpl.table_style == "grid":
                    table.style = d.styles["Table Grid"]
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table_borders(table)
                for i, col in enumerate(tb.columns):
                    c = table.rows[0].cells[i]
                    if tpl.table_style == "shaded":
                        shade(c, pal["fill"])
                    elif tpl.table_style == "booktabs":
                        cell_bottom(c, pal["accent"] if tpl.color else "000000")
                    c.paragraphs[0].add_run(col).bold = True
                for row in tb.rows:
                    cells = table.add_row().cells
                    for i, val in enumerate(row[:cols]):
                        cells[i].paragraphs[0].add_run(str(val))
                table_pt = tpl.body_pt - (1.5 if tpl.columns == 1 else 2)
                for row in table.rows:
                    for c in row.cells:
                        for par in c.paragraphs:
                            par.paragraph_format.space_after = Pt(1)
                            par.paragraph_format.line_spacing = 1.0
                            par.paragraph_format.first_line_indent = Cm(0)
                            for r in par.runs:
                                r.font.size = Pt(table_pt)
                n = d.add_paragraph()
                if tb.note:
                    run(n, "Note. ", size=tpl.body_pt - 2, italic=True)
                    run(n, tb.note, size=tpl.body_pt - 2)
            elif b.kind == "figure" and b.figure:
                img = a.project_dir / b.figure.path
                if img.exists():
                    from PIL import Image as PILImage

                    w_px, h_px = PILImage.open(img).size
                    fig_p = d.add_paragraph()
                    fig_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    fig_p.paragraph_format.keep_with_next = True
                    fig_p.paragraph_format.first_line_indent = Cm(0)
                    fig_p.add_run().add_picture(str(img), width=Cm(min(col_width_cm, 15.0, 11.5 * w_px / h_px)))
                    cap = d.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap.paragraph_format.first_line_indent = Cm(0)
                    run(cap, f"Figure {b.figure.number}. ", size=tpl.body_pt - 0.5, bold=True, color=accent)
                    run(cap, _clean_caption(b.figure.caption), size=tpl.body_pt - 0.5, italic=True)

    d.core_properties.title = a.title
    d.core_properties.author = ", ".join(a.authors) or "RAF"
    d.core_properties.keywords = a.keywords
    d.save(out)
    return out


# ====================================================================== PDF
def to_pdf(a: Article, out: Path, template: str | None = None, palette: str | None = None) -> Path:
    from xml.sax.saxutils import escape

    import matplotlib
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (BaseDocTemplate, Frame, FrameBreak, HRFlowable, Image, KeepTogether, NextPageTemplate, PageBreak,
                                    PageTemplate, Paragraph, Spacer, Table as RLTable, TableStyle)
    from reportlab.platypus.tableofcontents import TableOfContents

    tpl, pal = resolve_style(template, palette)
    font_dir = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
    fonts = {"serif": ("Times-Roman", "Times-Bold", "Times-Italic"), "sans": ("Helvetica", "Helvetica-Bold", "Helvetica-Oblique")}
    try:
        for name, file in (("RafSerif", "DejaVuSerif.ttf"), ("RafSerif-Bold", "DejaVuSerif-Bold.ttf"), ("RafSerif-Italic", "DejaVuSerif-Italic.ttf"),
                           ("RafSerif-BoldItalic", "DejaVuSerif-BoldItalic.ttf"),
                           ("RafSans", "DejaVuSans.ttf"), ("RafSans-Bold", "DejaVuSans-Bold.ttf"), ("RafSans-Oblique", "DejaVuSans-Oblique.ttf"),
                           ("RafSans-BoldOblique", "DejaVuSans-BoldOblique.ttf")):
            pdfmetrics.registerFont(TTFont(name, str(font_dir / file)))
        pdfmetrics.registerFontFamily("RafSerif", normal="RafSerif", bold="RafSerif-Bold", italic="RafSerif-Italic", boldItalic="RafSerif-BoldItalic")
        pdfmetrics.registerFontFamily("RafSans", normal="RafSans", bold="RafSans-Bold", italic="RafSans-Oblique", boldItalic="RafSans-BoldOblique")
        fonts = {"serif": ("RafSerif", "RafSerif-Bold", "RafSerif-Italic"), "sans": ("RafSans", "RafSans-Bold", "RafSans-Oblique")}
    except Exception:  # noqa: BLE001 - fall back to core fonts
        pass
    body_f = fonts["serif" if tpl.body_serif else "sans"]
    head_f = fonts["serif" if tpl.heading_serif else "sans"]
    accent = colors.HexColor("#" + pal["accent"])
    fill = colors.HexColor("#" + pal["fill"])
    rule = colors.HexColor("#" + pal["rule"])
    ink = colors.black
    muted = colors.HexColor("#5B5B66") if tpl.color else colors.black
    head_color = accent if tpl.color else ink
    pt, lead = tpl.body_pt, tpl.body_pt * tpl.leading
    acc_hex = pal["accent"] if tpl.color else "000000"

    S = {
        "body": ParagraphStyle("body", fontName=body_f[0], fontSize=pt, leading=lead, alignment=TA_JUSTIFY if tpl.justify else TA_LEFT,
                               spaceAfter=0 if tpl.leading >= 2 else pt * 0.6, firstLineIndent=tpl.indent_cm * cm),
        "h1": ParagraphStyle("h1", fontName=head_f[1], fontSize=pt + tpl.heading_scale, leading=(pt + tpl.heading_scale) * 1.3,
                             textColor=head_color, spaceBefore=pt * 1.4, spaceAfter=pt * 0.6, keepWithNext=1,
                             alignment=TA_CENTER if tpl.heading_center else TA_LEFT),
        "h2": ParagraphStyle("h2", fontName=head_f[1], fontSize=pt + max(0.5, tpl.heading_scale / 2), leading=(pt + 2) * 1.3,
                             textColor=head_color, spaceBefore=pt, spaceAfter=pt * 0.4, keepWithNext=1),
        "kicker": ParagraphStyle("kicker", fontName=head_f[1], fontSize=8, textColor=accent, spaceAfter=10, leading=10,
                                 alignment=TA_CENTER if tpl.title_align == "center" else TA_LEFT),
        "title": ParagraphStyle("title", fontName=head_f[1], fontSize=18 if tpl.columns == 2 else (pt if tpl.cover == "title_page" else 22),
                                leading=23 if tpl.columns == 2 else (lead if tpl.cover == "title_page" else 27), spaceAfter=6,
                                alignment=TA_CENTER if tpl.title_align == "center" else TA_LEFT),
        "subtitle": ParagraphStyle("subtitle", fontName=body_f[2], fontSize=12 if tpl.columns == 2 else 13, leading=17, textColor=muted, spaceAfter=10,
                                   alignment=TA_CENTER if tpl.title_align == "center" else TA_LEFT),
        "meta": ParagraphStyle("meta", fontName=body_f[0], fontSize=max(9, pt - 0.5), leading=max(9, pt - 0.5) * (tpl.leading if tpl.cover == "title_page" else 1.4),
                               textColor=muted, alignment=TA_CENTER if tpl.title_align == "center" else TA_LEFT),
        "abs": ParagraphStyle("abs", fontName=body_f[0], fontSize=pt - (0 if tpl.cover == "title_page" else 0.5), leading=lead * (1 if tpl.cover == "title_page" else 0.95),
                              alignment=TA_JUSTIFY if tpl.justify else TA_LEFT),
        "cap": ParagraphStyle("cap", fontName=body_f[0], fontSize=pt - 1, leading=(pt - 1) * 1.3, spaceBefore=8, spaceAfter=4, keepWithNext=1),
        "figcap": ParagraphStyle("figcap", fontName=body_f[0], fontSize=pt - 1, leading=(pt - 1) * 1.3, alignment=TA_CENTER, spaceBefore=4, spaceAfter=12),
        "cell": ParagraphStyle("cell", fontName=fonts["sans"][0] if tpl.color else body_f[0], fontSize=pt - (2.5 if tpl.columns == 1 else 2.2), leading=(pt - 2) * 1.25),
        "cellh": ParagraphStyle("cellh", fontName=fonts["sans"][1] if tpl.color else body_f[1], fontSize=pt - (2.5 if tpl.columns == 1 else 2.2), leading=(pt - 2) * 1.25),
        "note": ParagraphStyle("note", fontName=body_f[0], fontSize=pt - 2.5, leading=(pt - 2.5) * 1.3, textColor=muted, spaceAfter=10),
        "ref": ParagraphStyle("ref", fontName=body_f[0], fontSize=pt - (0 if tpl.cover == "title_page" else 1), leading=lead * (1 if tpl.cover == "title_page" else 0.9),
                              leftIndent=36 if tpl.cover == "title_page" else 18, firstLineIndent=-36 if tpl.cover == "title_page" else -18, spaceAfter=3),
        "toc1": ParagraphStyle("toc1", fontName=body_f[0], fontSize=pt, leading=pt * 1.6),
        "center_bold": ParagraphStyle("center_bold", fontName=body_f[1], fontSize=pt, leading=lead, alignment=TA_CENTER, spaceAfter=pt * 0.5),
    }

    class Doc(BaseDocTemplate):
        def afterFlowable(self, flowable):
            if isinstance(flowable, Paragraph) and flowable.style.name in {"h1", "h2"} and tpl.toc:
                level = 0 if flowable.style.name == "h1" else 1
                key = f"h{id(flowable)}"
                self.canv.bookmarkPage(key)
                self.canv.addOutlineEntry(flowable.getPlainText(), key, level=level)
                self.notify("TOCEntry", (level, flowable.getPlainText(), self.page, key))

    top, right, bottom, left = (v * cm for v in tpl.margins_cm)
    doc = Doc(str(out), pagesize=A4, leftMargin=left, rightMargin=right, topMargin=top, bottomMargin=bottom,
              title=a.title, author=", ".join(a.authors) or "RAF", subject="Research article", keywords=a.keywords)
    W, H = A4
    gutter = 0.6 * cm
    col_w = (doc.width - gutter) / 2
    short = a.title if len(a.title) < 95 else a.title[:92] + "…"

    def decorate(canvas, _doc, first_page: bool):
        canvas.saveState()
        if first_page and tpl.cover == "band":
            canvas.setFillColor(accent)
            canvas.rect(0, H - 0.45 * cm, W, 0.45 * cm, stroke=0, fill=1)
        canvas.setFont(head_f[0], 8 if tpl.cover != "title_page" else pt)
        canvas.setFillColor(muted)
        if tpl.cover == "title_page":
            canvas.drawRightString(W - right, H - top / 2, str(canvas.getPageNumber()))
        else:
            if tpl.running_header and not first_page:
                canvas.drawRightString(W - right, H - top + 0.55 * cm, short)
                canvas.setStrokeColor(rule)
                canvas.setLineWidth(0.5)
                canvas.line(left, H - top + 0.4 * cm, W - right, H - top + 0.4 * cm)
            canvas.drawCentredString(W / 2, bottom / 2, str(canvas.getPageNumber()))
        canvas.restoreState()

    def para(text: str, style: str) -> Paragraph:
        return Paragraph(escape(text), S[style])

    story: list = [NextPageTemplate("later")]
    content_w = col_w if tpl.columns == 2 else doc.width
    # ---- title block
    if tpl.cover == "title_page":
        story += [Spacer(1, 6 * cm), Paragraph(f"<b>{escape(a.title + (': ' + a.subtitle if a.subtitle else ''))}</b>", S["center_bold"]), Spacer(1, lead)]
        for line in [", ".join(a.authors), a.affiliation, a.date]:
            if line:
                story.append(Paragraph(escape(line), ParagraphStyle("tp", parent=S["meta"], textColor=ink)))
        story.append(PageBreak())
    else:
        story.append(Spacer(1, 0.3 * cm if tpl.columns == 2 else 1.0 * cm))
        if tpl.kicker:
            story.append(para(tpl.kicker, "kicker"))
        story.append(para(a.title, "title"))
        if a.subtitle:
            story.append(para(a.subtitle, "subtitle"))
        meta = []
        if a.authors:
            meta.append(f"<b>{escape(', '.join(a.authors))}</b>")
        if a.affiliation:
            meta.append(f"<i>{escape(a.affiliation)}</i>")
        meta.append(escape(a.date))
        story.append(Paragraph("<br/>".join(meta), S["meta"]))
        story.append(Spacer(1, 8))
        if tpl.cover in {"band", "rule"}:
            story.append(HRFlowable(width="100%", thickness=1.2, color=accent, spaceBefore=4, spaceAfter=10))

    # ---- abstract
    kw_label = "Index Terms—" if tpl.abstract_style == "ieee" else "Keywords:"
    if a.abstract:
        if tpl.abstract_style == "page":
            story += [Paragraph("<b>Abstract</b>", S["center_bold"]), Paragraph(escape(a.abstract), ParagraphStyle("absp", parent=S["abs"], alignment=TA_LEFT))]
            if a.keywords:
                story.append(Paragraph(f"<i>Keywords:</i> {escape(a.keywords.replace(';', ','))}", ParagraphStyle("kwp", parent=S["abs"], firstLineIndent=36)))
            story += [PageBreak(), Paragraph(f"<b>{escape(a.title + (': ' + a.subtitle if a.subtitle else ''))}</b>", S["center_bold"])]
        elif tpl.abstract_style == "box":
            inner = [Paragraph(f"<font color='#{acc_hex}'><b>Abstract</b></font>", S["abs"]), Spacer(1, 4), para(a.abstract, "abs")]
            if a.keywords:
                inner += [Spacer(1, 6), Paragraph(f"<b>Keywords:</b> <i>{escape(a.keywords)}</i>", S["abs"])]
            box = RLTable([[inner]], colWidths=[doc.width])
            box.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), fill), ("LINEBEFORE", (0, 0), (0, -1), 3, accent),
                                     ("LEFTPADDING", (0, 0), (-1, -1), 14), ("RIGHTPADDING", (0, 0), (-1, -1), 14),
                                     ("TOPPADDING", (0, 0), (-1, -1), 12), ("BOTTOMPADDING", (0, 0), (-1, -1), 12)]))
            story += [box, Spacer(1, 14)]
        elif tpl.abstract_style == "ieee":
            ieee = ParagraphStyle("ieee", parent=S["abs"], fontName=body_f[1], fontSize=pt - 0.5, leading=(pt - 0.5) * 1.25, alignment=TA_JUSTIFY)
            story.append(Paragraph(f"<i>Abstract</i>—{escape(a.abstract)}", ieee))
            if a.keywords:
                story += [Spacer(1, 4), Paragraph(f"<i>{kw_label}</i>{escape(a.keywords.replace(';', ','))}", ieee)]
        else:
            if tpl.abstract_style == "rules":
                story.append(HRFlowable(width="100%", thickness=0.8, color=accent, spaceAfter=6))
            story.append(Paragraph(f"<font color='#{acc_hex}'><b>{'ABSTRACT' if tpl.abstract_style == 'rules' else 'Abstract'}</b></font>",
                                   ParagraphStyle("abh", parent=S["abs"], fontName=head_f[1], fontSize=pt + (3 if tpl.abstract_style == "plain" else 0),
                                                  leading=(pt + 3) * 1.4, alignment=TA_CENTER if tpl.title_align == "center" else TA_LEFT, spaceAfter=4)))
            story.append(para(a.abstract, "abs"))
            if a.keywords:
                story += [Spacer(1, 5), Paragraph(f"<b>Keywords:</b> <i>{escape(a.keywords)}</i>", S["abs"])]
            if tpl.abstract_style == "rules":
                story.append(HRFlowable(width="100%", thickness=0.8, color=accent, spaceBefore=8, spaceAfter=6))
            story.append(Spacer(1, 10))
    if tpl.toc:
        toc = TableOfContents()
        toc.levelStyles = [S["toc1"], ParagraphStyle("toc2", parent=S["toc1"], leftIndent=16, fontSize=pt - 1)]
        story += [Paragraph(f"<font color='#{acc_hex}'><b>Contents</b></font>", S["body"]), toc, PageBreak()]
    # ---- page templates (the two-column title frame is sized to the measured title block)
    full = Frame(left, bottom, doc.width, doc.height, id="full")
    if tpl.columns == 2:
        head_h = 0.0
        for f in story[1:]:
            _, fh = f.wrap(doc.width - 12, doc.height)
            head_h += fh + f.getSpaceBefore() + f.getSpaceAfter()
        top_h = min(doc.height * 0.6, head_h + 24)
        first_frames = [Frame(left, bottom + doc.height - top_h, doc.width, top_h, id="top"),
                        Frame(left, bottom, col_w, doc.height - top_h, id="c1"),
                        Frame(left + col_w + gutter, bottom, col_w, doc.height - top_h, id="c2")]
        later_frames = [Frame(left, bottom, col_w, doc.height, id="l1"), Frame(left + col_w + gutter, bottom, col_w, doc.height, id="l2")]
        story.append(FrameBreak())
    else:
        first_frames, later_frames = [full], [Frame(left, bottom, doc.width, doc.height, id="full2")]
    doc.addPageTemplates([PageTemplate("first", first_frames, onPage=lambda c, d_: decorate(c, d_, True)),
                          PageTemplate("later", later_frames, onPage=lambda c, d_: decorate(c, d_, False))])

    # ---- body
    for s in a.sections:
        if s.key == "appendices" and tpl.columns == 1:
            story.append(PageBreak())
        story.append(para(heading_text(tpl, s.number, s.heading), "h1"))
        first = True
        for b in s.blocks:
            if b.kind == "p":
                style = S["body"]
                if first and tpl.key == "classic_journal":
                    style = ParagraphStyle("body0", parent=S["body"], firstLineIndent=0)
                story.append(Paragraph(escape(b.text), style))
                first = False
            elif b.kind == "subhead":
                story.append(para(b.text, "h2"))
            elif b.kind == "ref":
                story.append(para(b.text, "ref"))
            elif b.kind == "table" and b.table:
                tb = b.table
                cols = max(len(tb.columns), max((len(r) for r in tb.rows), default=1))
                data = [[Paragraph(escape(c), S["cellh"]) for c in tb.columns] + [""] * (cols - len(tb.columns))]
                data += [[Paragraph(escape(str(v)), S["cell"]) for v in r[:cols]] + [""] * (cols - len(r[:cols])) for r in tb.rows]
                weights = [max(4, min(40, max(len(str(r[i])) if i < len(r) else 0 for r in [tb.columns, *tb.rows]))) for i in range(cols)]
                widths = [content_w * w / sum(weights) for w in weights]
                t = RLTable(data, colWidths=widths, repeatRows=1)
                line_color = accent if tpl.color else ink
                cmds = [("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
                if tpl.table_style == "grid":
                    cmds += [("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#888888")), ("BACKGROUND", (0, 0), (-1, 0), fill if tpl.color else colors.HexColor("#EEEEEE"))]
                elif tpl.table_style == "booktabs":
                    cmds += [("LINEABOVE", (0, 0), (-1, 0), 1.1, line_color), ("LINEBELOW", (0, 0), (-1, 0), 0.6, line_color), ("LINEBELOW", (0, -1), (-1, -1), 1.1, line_color)]
                else:
                    cmds += [("BACKGROUND", (0, 0), (-1, 0), fill), ("LINEABOVE", (0, 0), (-1, 0), 1.1, accent), ("LINEBELOW", (0, 0), (-1, 0), 0.6, accent),
                             ("LINEBELOW", (0, -1), (-1, -1), 1.1, accent), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FAFAFC")])]
                t.setStyle(TableStyle(cmds))
                if tpl.cover == "title_page":
                    cap = Paragraph(f"<b>{table_label(tb)}</b><br/><i>{escape(tb.caption.rstrip('.'))}</i>", S["cap"])
                else:
                    cap = Paragraph(f"<font color='#{acc_hex}'><b>{table_label(tb)}.</b></font> <i>{escape(_clean_caption(tb.caption))}</i>", S["cap"])
                story += [cap, t, Paragraph(f"<i>Note.</i> {escape(tb.note)}", S["note"]) if tb.note else Spacer(1, 10)]
            elif b.kind == "figure" and b.figure:
                img_path = a.project_dir / b.figure.path
                if img_path.exists():
                    from PIL import Image as PILImage

                    w, h = PILImage.open(img_path).size
                    width = min(content_w, 15 * cm)
                    height = width * h / w
                    max_h = 11.5 * cm if tpl.columns == 1 else 8 * cm
                    if height > max_h:
                        height, width = max_h, max_h * w / h
                    story.append(KeepTogether([Image(str(img_path), width=width, height=height),
                                               Paragraph(f"<font color='#{acc_hex}'><b>Figure {b.figure.number}.</b></font> <i>{escape(_clean_caption(b.figure.caption))}</i>", S["figcap"])]))
    doc.multiBuild(story)
    return out


# ====================================================================== Markdown
def to_markdown(a: Article) -> str:
    lines = [f"# {a.title}"]
    if a.subtitle:
        lines.append(f"## {a.subtitle}")
    if a.authors:
        lines.append(f"**{', '.join(a.authors)}**  ")
    if a.affiliation:
        lines.append(f"*{a.affiliation}*  ")
    lines.append(a.date + "\n")
    if a.abstract:
        lines += ["## Abstract", a.abstract, ""]
    if a.keywords:
        lines += [f"**Keywords:** {a.keywords}", ""]
    for s in a.sections:
        lines.append(f"## {s.number + '. ' if s.number else ''}{s.heading}\n")
        for b in s.blocks:
            if b.kind == "p":
                lines.append(b.text + "\n")
            elif b.kind == "subhead":
                lines.append(f"### {b.text}\n")
            elif b.kind == "ref":
                lines.append(f"- {b.text}")
            elif b.kind == "table" and b.table:
                t = b.table
                lines.append(f"\n**{table_label(t)}.** *{_clean_caption(t.caption)}*\n")
                lines.append("| " + " | ".join(t.columns) + " |")
                lines.append("|" + "---|" * len(t.columns))
                for r in t.rows:
                    lines.append("| " + " | ".join(str(x).replace("|", "/") for x in r) + " |")
                if t.note:
                    lines.append(f"\n*Note.* {t.note}")
                lines.append("")
            elif b.kind == "figure" and b.figure:
                img = a.project_dir / b.figure.path
                if img.exists():
                    data = base64.b64encode(img.read_bytes()).decode()
                    lines.append(f"![Figure {b.figure.number}](data:image/png;base64,{data})\n")
                lines.append(f"**Figure {b.figure.number}.** *{_clean_caption(b.figure.caption)}*\n")
        lines.append("")
    return "\n".join(lines)


# ====================================================================== import
def import_document(p: Project, path: Path) -> dict[str, str]:
    """Split an edited DOCX / Markdown / PDF back into segments by matching its headings to segment titles."""
    from rapidfuzz import fuzz, process

    from .ingest import read_any

    doc = read_any(path)
    text = doc.text
    titles = {BLUEPRINTS[k].title.lower(): k for k in p.segments}
    aliases = {"literature review": "literature_review", "results": "results", "analysis": "results", "conclusion": "conclusion",
               "conclusions": "conclusion", "limitations": "limitations", "future research": "limitations", "bibliography": "references"}
    choices = {**titles, **{k: v for k, v in aliases.items() if v in p.segments}}

    lines = text.splitlines()
    marks: list[tuple[int, str]] = []
    for i, line in enumerate(lines):
        clean = re.sub(r"^(\d+(\.\d+)*|[IVXLC]+)\.?\s+", "", re.sub(r"^#+\s*", "", line.strip())).lower()
        if not clean or len(clean) > 60:
            continue
        match = process.extractOne(clean, list(choices), scorer=fuzz.ratio)
        if match and match[1] >= 88:
            marks.append((i, choices[match[0]]))
    out: dict[str, str] = {}
    for n, (start, key) in enumerate(marks):
        end = marks[n + 1][0] if n + 1 < len(marks) else len(lines)
        body = "\n".join(lines[start + 1 : end]).strip()
        body = re.sub(r"\n{3,}", "\n\n", body)
        if body and key not in out:
            out[key] = body
    return out
