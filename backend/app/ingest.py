"""Read PDFs, DOCX and Markdown/text files into clean text, sections and bibliographic metadata."""
from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

log = logging.getLogger("raf.ingest")

SUPPORTED = {".pdf": "pdf", ".docx": "docx", ".md": "md", ".markdown": "md", ".txt": "txt"}


@dataclass
class ParsedDoc:
    text: str
    pages: int = 0
    meta: dict = field(default_factory=dict)
    tables: list[list[list[str]]] = field(default_factory=list)
    headings: list[tuple[int, str]] = field(default_factory=list)


# ---------------------------------------------------------------- readers
def read_pdf(path: Path) -> ParsedDoc:
    try:
        import pymupdf

        doc = pymupdf.open(path)
        pages, tables = [], []
        for page in doc:
            pages.append(page.get_text("text"))
            try:
                for t in page.find_tables().tables[:3]:
                    data = [[(c or "").strip() for c in row] for row in t.extract()]
                    if len(data) > 1:
                        tables.append(data)
            except Exception:  # noqa: BLE001 - table finder is best-effort
                pass
        meta = {k: v for k, v in (doc.metadata or {}).items() if v}
        if len(doc):
            meta["visual_title"] = _largest_font_line(doc[0])
        return ParsedDoc(text=_join_pages(pages), pages=len(pages), meta=meta, tables=tables)
    except ImportError:
        pass
    except Exception as exc:  # noqa: BLE001
        log.warning("PyMuPDF failed on %s: %s, falling back", path.name, exc)

    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages = [(p.extract_text() or "") for p in reader.pages]
        meta = {k.strip("/").lower(): str(v) for k, v in (reader.metadata or {}).items() if v}
        doc = ParsedDoc(text=_join_pages(pages), pages=len(pages), meta=meta)
    except Exception as exc:  # noqa: BLE001
        log.warning("pypdf failed on %s: %s", path.name, exc)
        doc = ParsedDoc(text="")

    if len(doc.text) < 200:
        try:
            import pdfplumber

            with pdfplumber.open(path) as pdf:
                pages = [(p.extract_text() or "") for p in pdf.pages]
                doc.tables = [t for p in pdf.pages[:30] for t in (p.extract_tables() or [])]
                doc.text, doc.pages = _join_pages(pages), len(pages)
        except Exception as exc:  # noqa: BLE001
            log.warning("pdfplumber failed on %s: %s", path.name, exc)
    return doc


def _largest_font_line(page) -> str:
    """The title of a paper is almost always the largest text on its first page."""
    try:
        lines = []
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                text = "".join(s["text"] for s in line["spans"]).strip()
                size = max((s["size"] for s in line["spans"]), default=0)
                if len(text) > 3:
                    lines.append((round(size, 1), line["bbox"][1], text))
        if not lines:
            return ""
        top = max(l[0] for l in lines)
        title_lines = [l for l in sorted(lines, key=lambda l: l[1]) if l[0] == top]
        title = " ".join(l[2] for l in title_lines[:4])
        return title if 10 <= len(title) <= 300 else ""
    except Exception:  # noqa: BLE001
        return ""


def read_docx(path: Path) -> ParsedDoc:
    import docx

    d = docx.Document(str(path))
    parts, headings = [], []
    for para in d.paragraphs:
        t = para.text.strip()
        if not t:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading") or style == "title":
            level = int(re.sub(r"\D", "", style) or 1)
            headings.append((level, t))
            parts.append(f"\n{'#' * level} {t}\n")
        else:
            parts.append(t)
    tables = [[[c.text.strip() for c in row.cells] for row in tbl.rows] for tbl in d.tables]
    cp = d.core_properties
    meta = {"title": cp.title or "", "author": cp.author or "", "year": str(cp.created.year) if cp.created else ""}
    return ParsedDoc(text="\n\n".join(parts), pages=0, meta={k: v for k, v in meta.items() if v}, tables=tables, headings=headings)


def read_markdown(path: Path) -> ParsedDoc:
    text = path.read_text(encoding="utf-8", errors="ignore")
    headings = [(len(m.group(1)), m.group(2).strip()) for m in re.finditer(r"^(#{1,6})\s+(.+)$", text, re.M)]
    tables = []
    for block in re.findall(r"((?:^\|.*\|\s*$\n?){2,})", text, re.M):
        rows = [[c.strip() for c in line.strip().strip("|").split("|")] for line in block.strip().splitlines()]
        rows = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
        if rows:
            tables.append(rows)
    return ParsedDoc(text=text, meta={"title": headings[0][1]} if headings else {}, tables=tables, headings=headings)


def read_any(path: Path) -> ParsedDoc:
    kind = SUPPORTED.get(path.suffix.lower())
    if kind == "pdf":
        return read_pdf(path)
    if kind == "docx":
        return read_docx(path)
    if kind in {"md", "txt"}:
        return read_markdown(path)
    raise ValueError(f"Unsupported file type: {path.suffix}")


# ---------------------------------------------------------------- cleaning
def _join_pages(pages: list[str]) -> str:
    text = "\n\n".join(pages)
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)            # de-hyphenate line breaks
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"(?<![.\n:;])\n(?!\n)", " ", text)        # unwrap soft line breaks
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def strip_back_matter(text: str) -> str:
    """Drop the reference list of a source article so it does not pollute retrieval."""
    m = re.search(r"\n\s*(references|bibliography|works cited|literature cited)\s*\n", text, re.I)
    if m and m.start() > len(text) * 0.5:
        return text[: m.start()]
    return text


# ---------------------------------------------------------------- metadata
_DOI = re.compile(r"\b(10\.\d{4,9}/[^\s\"<>]+)", re.I)
_YEAR = re.compile(r"\b(19[5-9]\d|20[0-4]\d)\b")


_PLACEHOLDER_META = re.compile(r"^\(?(anonymous|untitled|unknown|none|user|admin|author)\)?$|microsoft word|^document\d*$|\.(pdf|docx?|tex)$", re.I)


def heuristic_metadata(doc: ParsedDoc, filename: str) -> dict:
    head = doc.text[:4000]
    doi = _DOI.search(doc.text[:20000])
    years = _YEAR.findall(head)
    title = (doc.meta.get("title") or "").strip()
    if not title or len(title) < 8 or _PLACEHOLDER_META.search(title):
        title = doc.meta.get("visual_title", "")
    if not title:
        lines = [l.strip() for l in re.split(r"\n|(?<=[a-z])\s{2,}", head) if 15 < len(l.strip()) < 220]
        title = lines[0] if lines else Path(filename).stem.replace("_", " ")
    author = doc.meta.get("author", "")
    if _PLACEHOLDER_META.search(author.strip()):
        author = ""
    return {
        "title": title[:300],
        "authors": [a.strip() for a in re.split(r";|,| and ", author) if a.strip()][:8] if author else [],
        "year": years[0] if years else "",
        "doi": doi.group(1).rstrip(".,;)") if doi else "",
        "venue": "",
    }


def llm_metadata(doc: ParsedDoc, fallback: dict) -> dict:
    """Ask the SLM to read the first page and return clean citation metadata."""
    from . import llm

    system = "You extract bibliographic metadata from the first page of a scholarly document. Reply with JSON only."
    prompt = (
        "Return JSON with keys: title (string), authors (list of 'Surname, Initials' strings), "
        "year (4-digit string or empty), venue (journal/conference/publisher or empty), doi (or empty).\n"
        "Only use what is printed on the page; leave a field empty when unsure.\n\n"
        f"FIRST PAGE:\n{doc.text[:3500]}"
    )
    data = llm.generate_json(system, prompt, default={}, max_tokens=400)
    if not isinstance(data, dict):
        return fallback
    out = dict(fallback)
    for key in ("title", "year", "venue", "doi"):
        val = data.get(key)
        if isinstance(val, str) and val.strip():
            out[key] = val.strip()
    if isinstance(data.get("authors"), list) and data["authors"]:
        out["authors"] = [str(a).strip() for a in data["authors"] if str(a).strip()][:12]
    if out.get("year") and not _YEAR.fullmatch(str(out["year"])):
        out["year"] = fallback.get("year", "")
    return out


# ---------------------------------------------------------------- chunking
def chunk_text(text: str, target_words: int = 220, overlap_words: int = 40) -> list[str]:
    """Paragraph-aware sliding window chunker."""
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if len(p.strip()) > 40]
    chunks: list[str] = []
    buf: list[str] = []
    count = 0
    for para in paragraphs:
        words = para.split()
        if len(words) > target_words * 1.6:  # split very long paragraphs on sentences
            for sent_chunk in _split_long(para, target_words):
                chunks.append(sent_chunk)
            continue
        buf.append(para)
        count += len(words)
        if count >= target_words:
            chunks.append(" ".join(buf))
            tail = " ".join(" ".join(buf).split()[-overlap_words:])
            buf, count = [tail], overlap_words
    if buf and count > overlap_words:
        chunks.append(" ".join(buf))
    return [c for c in chunks if len(c.split()) > 25]


def _split_long(para: str, target: int) -> list[str]:
    sentences = split_sentences(para)
    out, buf, n = [], [], 0
    for s in sentences:
        buf.append(s)
        n += len(s.split())
        if n >= target:
            out.append(" ".join(buf))
            buf, n = buf[-1:], len(buf[-1].split())
    if buf:
        out.append(" ".join(buf))
    return out


_SENT = re.compile(r"(?<=[.!?])\s+(?=[A-Z(\[\"'])")


def split_sentences(text: str) -> list[str]:
    protected = re.sub(r"\b(e\.g|i\.e|et al|etc|vs|Fig|No|Dr|pp|cf)\.", lambda m: m.group(0).replace(".", "§"), text)
    return [s.replace("§", ".").strip() for s in _SENT.split(protected) if s.strip()]
